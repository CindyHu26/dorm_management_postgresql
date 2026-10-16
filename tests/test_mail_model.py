import tempfile
import unittest
import io
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from docx import Document
from openpyxl import Workbook, load_workbook

from data_models import mail_model
from views import mail_view


class MailDocumentTests(unittest.TestCase):
    def _items(self, count):
        return [
            {
                "sequence_no": index,
                "recipient_name_snapshot": f"收件人{index}",
                "recipient_postal_code_snapshot": f"郵遞區號{index}",
                "recipient_address_snapshot": f"地址{index}",
                "item_description": f"文件{index}",
                "tracking_no": None,
                "shipping_fee": None,
                "billing_employer_names": "政泰、台彰",
                "notes": "請附回郵信封",
            }
            for index in range(1, count + 1)
        ]

    def test_admin_code_hash_can_be_verified(self):
        encoded = mail_model._hash_admin_code("2468", iterations=1000)
        algorithm, iterations, salt_text, _ = encoded.split("$", 3)
        self.assertEqual("pbkdf2_sha256", algorithm)
        import base64

        recreated = mail_model._hash_admin_code(
            "2468",
            base64.urlsafe_b64decode(salt_text.encode("ascii")),
            int(iterations),
        )
        self.assertEqual(encoded, recreated)

    def test_audit_only_writes_important_events(self):
        cursor = MagicMock()

        mail_model._audit(
            cursor,
            "MAIL_RECORD",
            5,
            "CANCEL",
            1,
            after={"reason": "重複建立"},
        )
        cursor.execute.assert_called_once()

        cursor.reset_mock()
        mail_model._audit(
            cursor,
            "MAIL_RECORD",
            6,
            "CREATE",
            1,
            after={"mail_no": "MAIL-6"},
        )
        cursor.execute.assert_not_called()

    def test_audit_ignores_session_and_routine_save_events(self):
        cursor = MagicMock()

        for entity_type, action in (
            ("ADMIN_SESSION", "LOGIN"),
            ("MAIL_RECORD", "SAVE_PLANNED_POSTAGE"),
            ("MAIL_BATCH", "SAVE_TRACKING_NUMBERS"),
            ("MAIL_USER", "BATCH_SAVE"),
            ("MAIL_AGENCY", "BATCH_SAVE"),
            ("MAIL_RECIPIENT_LIST", "UPDATE"),
        ):
            mail_model._audit(cursor, entity_type, 1, action, 1)

        cursor.execute.assert_not_called()

    def test_shipping_fee_parser_accepts_editor_empty_and_display_values(self):
        self.assertIsNone(mail_model._parse_shipping_fee(None))
        self.assertIsNone(mail_model._parse_shipping_fee(""))
        self.assertIsNone(mail_model._parse_shipping_fee(float("nan")))
        self.assertEqual("28", str(mail_model._parse_shipping_fee(28)))
        self.assertEqual("28", str(mail_model._parse_shipping_fee([28])))
        self.assertEqual("28", str(mail_model._parse_shipping_fee("[28]")))
        self.assertEqual("1200", str(mail_model._parse_shipping_fee("NT$ 1,200")))
        with self.assertRaisesRegex(ValueError, "單一數字"):
            mail_model._parse_shipping_fee([28, 30])
        with self.assertRaisesRegex(ValueError, "不接受小數"):
            mail_model._parse_shipping_fee("28.5")

    def test_billing_names_can_fill_description_for_single_or_multiple_employers(self):
        state = SimpleNamespace()
        state.get = lambda key, default=None: getattr(state, key, default)
        with patch.object(mail_view.st, "session_state", state):
            mail_view._apply_billing_names_to_description(["政泰"])
            self.assertEqual("政泰", state.mail_form_description)

            mail_view._apply_billing_names_to_description(["政泰", "台彰", "佰藝", "弘霖玻璃"])
            self.assertEqual("政泰、台彰、佰藝、弘霖玻璃", state.mail_form_description)

            state.mail_form_description_pick = "稅務文件"
            mail_view._append_common_description()
            self.assertEqual("稅務文件-政泰、台彰、佰藝、弘霖玻璃", state.mail_form_description)

    def test_common_descriptions_include_previous_and_current_roc_tax_months(self):
        descriptions = mail_view._common_item_descriptions(date(2026, 10, 14))
        self.assertEqual(["115年9月稅", "115年10月稅"], descriptions[-2:])

    def test_common_tax_months_cross_roc_year_boundary(self):
        descriptions = mail_view._common_item_descriptions(date(2027, 1, 2))
        self.assertEqual(["115年12月稅", "116年1月稅"], descriptions[-2:])

    def test_pending_postage_values_uses_typed_sequence_and_postage(self):
        rows = [
            {"id": 11, "順序": 20, "郵資": [36]},
            {"id": 12, "順序": 5, "郵資": None},
            {"id": 13, "順序": 10, "郵資": "[28]"},
        ]

        record_ids, item_fees = mail_view._pending_postage_values(rows)

        self.assertEqual([12, 13, 11], record_ids)
        self.assertEqual({11: 36, 12: None, 13: 28}, item_fees)

    def test_pending_postage_values_rejects_duplicate_sequence(self):
        rows = [
            {"id": 11, "順序": 1, "郵資": None},
            {"id": 12, "順序": 1, "郵資": None},
        ]

        with self.assertRaisesRegex(ValueError, "順序 1 重複"):
            mail_view._pending_postage_values(rows)

    def test_pending_draft_values_allows_blank_and_duplicate_sequences(self):
        rows = [
            {"id": 11, "順序": "[1]", "常用郵資": "60", "郵資": 999},
            {"id": 12, "順序": [None], "郵資": None},
            {"id": 13, "順序": 1, "郵資": 36},
        ]

        item_sequences, item_fees = mail_view._pending_draft_values(rows)

        self.assertEqual({11: 1, 12: None, 13: 1}, item_sequences)
        self.assertEqual({11: 60, 12: None, 13: 36}, item_fees)

    def test_pending_draft_values_treats_32_as_custom_postage(self):
        item_sequences, item_fees = mail_view._pending_draft_values(
            [{"id": 11, "順序": 1, "常用郵資": "32", "郵資": 32}]
        )

        self.assertEqual({11: 1}, item_sequences)
        self.assertEqual({11: 32}, item_fees)

    def test_postal_sort_key_places_numeric_codes_first(self):
        rows = [
            {"id": 3, "recipient_postal_code_snapshot": "", "recipient_name_snapshot": "無區號"},
            {"id": 2, "recipient_postal_code_snapshot": "407", "recipient_name_snapshot": "台中"},
            {"id": 1, "recipient_postal_code_snapshot": "100", "recipient_name_snapshot": "台北"},
        ]

        ordered = sorted(rows, key=mail_view._postal_sort_key)

        self.assertEqual([1, 2, 3], [row["id"] for row in ordered])

    def test_postal_sort_key_uses_first_three_digits_for_delivery_area(self):
        rows = [
            {"id": 2, "recipient_postal_code_snapshot": "505", "recipient_name_snapshot": "鹿港"},
            {"id": 1, "recipient_postal_code_snapshot": "50001", "recipient_name_snapshot": "彰化"},
        ]

        ordered = sorted(rows, key=mail_view._postal_sort_key)

        self.assertEqual([1, 2], [row["id"] for row in ordered])

    def test_postal_sort_key_infers_missing_code_without_changing_record(self):
        row = {
            "id": 1,
            "recipient_postal_code_snapshot": None,
            "recipient_name_snapshot": "測試收件人",
            "recipient_address_snapshot": "彰化縣鹿港鎮民權路163號3樓",
        }

        sort_key = mail_view._postal_sort_key(row)

        self.assertEqual(505, sort_key[1])
        self.assertIsNone(row["recipient_postal_code_snapshot"])

    def test_saved_postal_code_takes_priority_over_inferred_default(self):
        row = {
            "recipient_postal_code_snapshot": "505001",
            "recipient_address_snapshot": "彰化縣鹿港鎮民權路163號3樓",
        }

        self.assertEqual("505001", mail_model.resolved_recipient_postal_code(row))

    def test_postage_summary_groups_fees_and_builds_equation(self):
        counts, equation, total = mail_view._postage_summary({1: 32, 2: 44, 3: 32})

        self.assertEqual({32: 2, 44: 1}, dict(counts))
        self.assertEqual("32 元 × 2 封 ＋ 44 元 × 1 封", equation)
        self.assertEqual(108, total)

    def test_mail_service_summary_groups_pending_envelopes(self):
        rows = [
            {"mail_service": "掛號"},
            {"mail_service": "平信"},
            {"mail_service": "掛號"},
            {"mail_service": "掛號"},
        ]

        self.assertEqual(
            "平信 1 封、掛號 3 封、共 4 封",
            mail_view._mail_service_summary(rows),
        )

    def test_mail_workbench_meta_hides_blank_notes(self):
        row = {
            "recipient_address_snapshot": "彰化市中山路 1 號",
            "mail_service": "掛號",
            "notes": "",
        }

        meta = mail_view._mail_workbench_meta(row)

        self.assertIn("彰化市中山路 1 號", meta)
        self.assertIn("掛號", meta)
        self.assertNotIn("種類：", meta)
        self.assertNotIn("備註", meta)

    def test_mail_queue_keeps_saved_order_when_pending_order_changes(self):
        pending_by_id = {2: {"id": 2}, 3: {"id": 3}, 1: {"id": 1}}

        queue_ids = mail_view._reconcile_mail_queue([3, 1, 2], pending_by_id)

        self.assertEqual([3, 1, 2], queue_ids)

    def test_mail_queue_only_removes_unavailable_or_duplicate_ids(self):
        pending_by_id = {1: {"id": 1}, 2: {"id": 2}}

        queue_ids = mail_view._reconcile_mail_queue([2, 99, 2, "1"], pending_by_id)

        self.assertEqual([2, 1], queue_ids)

    def test_mail_record_option_uses_recipient_details_without_mail_number(self):
        row = {
            "mail_no": "MAIL-20261016-00001",
            "recipient_postal_code_snapshot": "50001",
            "recipient_name_snapshot": "彰化縣政府勞工處",
            "recipient_address_snapshot": "彰化市中興路 100 號",
            "item_description": "逃跑報備",
            "mail_service": "掛號",
            "notes": "金瑞瑩",
            "status": "READY",
        }

        label = mail_view._mail_record_option_label(row, include_status=True)

        self.assertNotIn("MAIL-20261016-00001", label)
        self.assertIn("【50001】彰化縣政府勞工處", label)
        self.assertIn("彰化市中興路 100 號", label)
        self.assertIn("逃跑報備", label)
        self.assertIn("備註：金瑞瑩", label)
        self.assertIn("待寄出", label)

    def test_mail_record_option_can_include_taipei_created_time(self):
        row = {
            "recipient_name_snapshot": "彰化縣政府勞工處",
            "created_at": datetime(2026, 10, 16, 1, 2, 3, tzinfo=timezone.utc),
        }

        label = mail_view._mail_record_option_label(row, include_created_at=True)

        self.assertIn("建立時間：2026-10-16 09:02:03", label)

    def test_save_pending_postage_persists_partial_sequence_progress(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [
            {"id": 11, "status": "READY"},
            {"id": 12, "status": "READY"},
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message = mail_model.save_pending_postage(
                {11: 28, 12: None},
                user_id=7,
                item_sequences={11: 1, 12: None},
            )

        self.assertTrue(success, message)
        update_params = [
            call.args[1]
            for call in cursor.execute.call_args_list
            if "SET planned_sequence = %s" in call.args[0]
        ]
        self.assertEqual([(1, 28, 11), (None, None, 12)], update_params)
        connection.commit.assert_called_once_with()

    def test_recipient_list_excel_import_accepts_supported_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["收件人名稱", "郵遞區號", "地址", "部門", "聯絡人", "備註"])
        sheet.append(["彰化縣政府", "500", "彰化市中山路1號", "勞工處", "林小姐", "每月寄送"])
        output = io.BytesIO()
        workbook.save(output)

        items = mail_model.parse_recipient_list_excel(output.getvalue())

        self.assertEqual(1, len(items))
        self.assertEqual("彰化縣政府", items[0]["recipient_name"])
        self.assertEqual("500", items[0]["recipient_postal_code"])
        self.assertEqual("勞工處", items[0]["recipient_department"])
        self.assertEqual("CUSTOM", items[0]["recipient_type"])

    def test_recipient_list_excel_import_requires_name_and_address_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["公司", "電話"])
        sheet.append(["測試公司", "04-1234567"])
        output = io.BytesIO()
        workbook.save(output)

        with self.assertRaisesRegex(ValueError, "收件人名稱.*地址"):
            mail_model.parse_recipient_list_excel(output.getvalue())

    def test_recipient_list_rejects_duplicate_people(self):
        items, error = mail_model._normalize_recipient_list_items(
            [
                {"recipient_name": "測試公司", "recipient_address": "彰化市中山路1號"},
                {"recipient_name": " 測試公司 ", "recipient_address": "彰化市中山路1號"},
            ]
        )
        self.assertIsNone(items)
        self.assertIn("重複", error)

    def test_imported_recipients_add_only_missing_shared_agencies(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [
            {
                "id": 1,
                "agency_name": "既有機構",
                "department_name": "人事部",
                "address": "既有地址",
            }
        ]
        items = [
            {
                "recipient_name": "既有機構",
                "recipient_department": "人事部",
                "recipient_address": "既有地址",
            },
            {
                "recipient_name": "新機構",
                "recipient_department": "收發室",
                "recipient_postal_code": "500",
                "recipient_address": "新地址",
                "recipient_contact": "林小姐",
            },
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.add_recipient_items_to_agencies(items, 7)

        self.assertTrue(success, message)
        agency_inserts = [
            call.args[1]
            for call in cursor.execute.call_args_list
            if "INSERT INTO mail_government_agencies" in call.args[0]
        ]
        self.assertEqual(
            [("新機構", "收發室", "林小姐", "500", "新地址")],
            agency_inserts,
        )
        self.assertIn("新增 1 筆", message)
        self.assertIn("略過既有 1 筆", message)
        connection.commit.assert_called_once_with()
        audit.assert_called_once()

    def test_sequential_tracking_numbers_preserve_width_and_suffix(self):
        numbers = mail_model.generate_sequential_tracking_numbers("0008", "12345", 3)
        self.assertEqual(["000812345", "000912345", "001012345"], numbers)

    def test_sequential_tracking_numbers_reject_overflow(self):
        with self.assertRaisesRegex(ValueError, "超過 2 碼"):
            mail_model.generate_sequential_tracking_numbers("99", "12345", 2)

    def test_save_batch_tracking_numbers_updates_items_in_sequence(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {"status": "CREATED"}
        cursor.fetchall.return_value = [{"id": 21}, {"id": 22}]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.save_batch_tracking_numbers(
                7, 1, ["000812345", "000912345"]
            )

        self.assertTrue(success, message)
        update_calls = [
            call.args[1]
            for call in cursor.execute.call_args_list
            if "SET tracking_no = %s" in call.args[0]
        ]
        self.assertEqual([("000812345", 21), ("000912345", 22)], update_calls)
        connection.commit.assert_called_once_with()
        audit.assert_called_once()

    def test_mail_user_editor_saves_changes_and_new_rows(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [{"id": 1}]
        rows = [
            {"id": 1, "啟用": True, "姓名": "王小明", "排序": 2},
            {"id": None, "啟用": True, "姓名": "陳小華", "排序": 1},
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.save_mail_users(rows, actor_user_id=1)

        self.assertTrue(success, message)
        self.assertIn("新增 1 筆", message)
        self.assertIn("更新 1 筆", message)
        connection.commit.assert_called_once_with()
        audit.assert_called_once()

    def test_mail_user_editor_requires_one_active_user(self):
        success, message = mail_model.save_mail_users(
            [{"id": 1, "啟用": False, "姓名": "王小明", "排序": 0}],
            actor_user_id=1,
        )
        self.assertFalse(success)
        self.assertIn("至少需要保留一位", message)

    def test_government_agency_editor_only_updates_changed_rows(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [
            {
                "id": 1,
                "is_active": True,
                "agency_name": "勞工局",
                "department_name": "外勞科",
                "postal_code": "408221",
                "address": "台中市南屯區精科路26號",
            },
            {
                "id": 2,
                "is_active": True,
                "agency_name": "移民署",
                "department_name": None,
                "postal_code": "100",
                "address": "台北市中正區廣州街15號",
            },
        ]
        rows = [
            {
                "id": 1,
                "啟用": True,
                "機關名稱": "勞工局",
                "部門": "外勞科",
                "郵遞區號": "408221",
                "地址": "台中市南屯區精科路26號",
            },
            {
                "id": 2,
                "啟用": True,
                "機關名稱": "移民署",
                "部門": "服務站",
                "郵遞區號": "100",
                "地址": "台北市中正區廣州街15號",
            },
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.save_government_agencies(rows, actor_user_id=1)

        self.assertTrue(success, message)
        self.assertIn("新增 0 筆、更新 1 筆、停用 0 筆", message)
        update_calls = [
            call for call in cursor.execute.call_args_list
            if "UPDATE mail_government_agencies SET" in call.args[0]
        ]
        self.assertEqual(1, len(update_calls))
        self.assertEqual(2, update_calls[0].args[1][-1])
        self.assertEqual(1, audit.call_args.kwargs["after"]["updated"])

    def test_government_agency_editor_reports_no_change(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [
            {
                "id": 1,
                "is_active": True,
                "agency_name": "勞工局",
                "department_name": None,
                "postal_code": "408221",
                "address": "台中市南屯區精科路26號",
            }
        ]
        rows = [
            {
                "id": 1,
                "啟用": True,
                "機關名稱": "勞工局",
                "部門": None,
                "郵遞區號": "408221",
                "地址": "台中市南屯區精科路26號",
            }
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message = mail_model.save_government_agencies(rows, actor_user_id=1)

        self.assertTrue(success, message)
        self.assertIn("新增 0 筆、更新 0 筆、停用 0 筆", message)
        self.assertFalse(
            any(
                "UPDATE mail_government_agencies SET" in call.args[0]
                for call in cursor.execute.call_args_list
            )
        )
        connection.commit.assert_called_once_with()

    def test_government_agency_editor_saves_employer_relations(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.side_effect = [
            [
                {
                    "id": 1,
                    "is_active": True,
                    "agency_name": "安心會計師事務所",
                    "department_name": None,
                    "contact_person": "林小姐",
                    "postal_code": "500",
                    "address": "彰化市中山路1號",
                }
            ],
            [],
        ]
        rows = [
            {
                "id": 1,
                "啟用": True,
                "機關名稱": "安心會計師事務所",
                "部門": None,
                "聯絡人": "林小姐",
                "郵遞區號": "500",
                "地址": "彰化市中山路1號",
                "關聯雇主": ["A001｜東方公司"],
            }
        ]
        employer_options = [
            {
                "key": "CODE:a001",
                "employer_id": "A001",
                "employer_name": "東方公司",
                "label": "A001｜東方公司",
            }
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "get_employer_relation_options", return_value=employer_options),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.save_government_agencies(rows, actor_user_id=1)

        self.assertTrue(success, message)
        self.assertIn("更新 1 筆", message)
        relation_insert = next(
            call
            for call in cursor.execute.call_args_list
            if "INSERT INTO mail_government_agency_employers" in call.args[0]
        )
        self.assertEqual((1, "CODE:a001", "A001", "東方公司"), relation_insert.args[1])
        self.assertEqual(1, audit.call_args.kwargs["after"]["updated"])
        connection.commit.assert_called_once_with()

    def test_employer_recipient_includes_workplace_postal_code(self):
        employer_rows = [
            {
                "id": 7,
                "employer_id": "EMP007",
                "employer_name": "測試公司",
                "address": "彰化縣福興鄉福工路2之1號",
                "postal_code": "506",
            }
        ]
        with patch.object(mail_model, "_fetch_all", side_effect=[employer_rows, []]):
            options = mail_model.get_recipient_options()

        self.assertEqual("506", options[0]["postal_code"])
        self.assertEqual("E:7", options[0]["key"])

    def test_agency_recipient_includes_contact_and_related_employers(self):
        agency_rows = [
            {
                "id": 3,
                "agency_name": "安心會計師事務所",
                "department_name": None,
                "contact_person": "林小姐",
                "postal_code": "500",
                "address": "彰化市中山路1號",
                "related_employers": ["A001｜東方公司", "A018｜永盛公司"],
                "related_employer_keys": ["CODE:a001", "CODE:a018"],
            }
        ]
        with patch.object(mail_model, "_fetch_all", side_effect=[[], agency_rows]):
            options = mail_model.get_recipient_options()

        self.assertEqual("林小姐", options[0]["contact"])
        self.assertEqual(["CODE:a001", "CODE:a018"], options[0]["related_employer_keys"])
        self.assertIn("A001｜東方公司", options[0]["label"])
        self.assertIn("A018｜永盛公司", options[0]["label"])
        self.assertIn("安心會計師事務所｜彰化市中山路1號｜林小姐", options[0]["label"])

    def test_recipient_options_include_current_users_personal_list_members(self):
        employer_rows = [
            {
                "id": 7,
                "employer_id": "EMP007",
                "employer_name": "既有客戶",
                "address": "客戶地址",
                "postal_code": "500",
            }
        ]
        personal_rows = [
            {
                "id": 31,
                "recipient_type": "CUSTOM",
                "employer_workplace_id": None,
                "government_agency_id": None,
                "recipient_name": "名單內收件人",
                "recipient_department": "收發室",
                "recipient_postal_code": "500",
                "recipient_address": "彰化市測試路1號",
                "recipient_contact": "林小姐",
                "billing_employers": [],
                "list_name": "每月寄送",
            }
        ]
        with patch.object(mail_model, "_fetch_all", side_effect=[employer_rows, [], personal_rows]) as fetch_all:
            options = mail_model.get_recipient_options(user_id=7)

        self.assertEqual(2, len(options))
        self.assertEqual("L:31", options[0]["key"])
        self.assertEqual("名單內收件人", options[0]["name"])
        self.assertIn("我的名單｜每月寄送", options[0]["label"])
        self.assertEqual("E:7", options[1]["key"])
        self.assertEqual((7,), fetch_all.call_args_list[2].args[1])

    def test_company_recipient_shortcut_contains_fixed_company_details(self):
        option = mail_model.get_company_recipient_option()

        self.assertEqual("COMPANY", option["key"])
        self.assertEqual("長青國際集團", option["name"])
        self.assertEqual("505", option["postal_code"])
        self.assertEqual("彰化縣鹿港鎮民權路163號3樓", option["address"])

    def test_employer_relation_options_deduplicate_multiple_workplaces(self):
        workplaces = [
            {"employer_id": "A001", "employer_name": "東方公司"},
            {"employer_id": "A001", "employer_name": "東方公司"},
            {"employer_id": None, "employer_name": "無代號公司"},
        ]
        with patch.object(mail_model, "_fetch_all", return_value=workplaces):
            options = mail_model.get_employer_relation_options()

        self.assertEqual(2, len(options))
        self.assertEqual("A001｜東方公司", options[0]["label"])
        self.assertEqual("無代號公司", options[1]["label"])

    def test_bulk_workbook_uses_template_cells_and_adds_pages(self):
        batch = {
            "batch_no": "POST-TEST-00001",
            "submission_mode": "BULK_FORM",
            "template_filename": "郵局大宗掛號單.xlsx",
        }
        with patch.object(mail_model, "get_batch", return_value=(batch, self._items(41))):
            data, filename = mail_model.generate_bulk_workbook(1)

        self.assertEqual("郵局大宗掛號單_POST-TEST-00001.xlsx", filename)
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "result.xlsx"
            output.write_bytes(data)
            workbook = load_workbook(output)
        self.assertEqual(["P1", "P2"], workbook.sheetnames)
        self.assertEqual("收件人1", workbook["P1"]["D5"].value)
        self.assertEqual("郵遞區號1  文件1  請附回郵信封", workbook["P1"]["H5"].value)
        self.assertNotIn("地址1", workbook["P1"]["H5"].value)
        self.assertEqual("收件人41", workbook["P2"]["D5"].value)
        self.assertIsNone(workbook["P2"]["D6"].value)
        self.assertEqual("共", workbook["P1"]["F45"].value)
        self.assertEqual(40, workbook["P1"]["G45"].value)
        self.assertEqual("件", workbook["P1"]["I45"].value)
        self.assertEqual(1, workbook["P2"]["G45"].value)
        self.assertIn("$A$1:$J$45", str(workbook["P2"].print_area))
        self.assertIn("$1:$4", str(workbook["P2"].print_title_rows))

    def test_b4_envelope_fills_frontend_labels(self):
        record = {
            "mail_no": "MAIL-TEST-1",
            "recipient_postal_code_snapshot": "408221",
            "recipient_name_snapshot": "台中市政府勞工局",
            "recipient_department_snapshot": "入國通報",
            "recipient_contact_snapshot": "林小姐",
            "recipient_address_snapshot": "台中市南屯區精科路26號",
            "sender_name_snapshot": "王小明",
            "item_description": "D02-阮氏水",
            "mail_service": "限時掛號",
        }
        with patch.object(mail_model, "get_mail_record", return_value=record):
            data, filename = mail_model.generate_envelope(1, "B4.docx")

        document = Document(io.BytesIO(data))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        self.assertEqual(
            [
                "限時掛號",
                "408221",
                "台中市南屯區精科路26號",
                "台中市政府勞工局   收",
                "入國通報：林小姐",
                "D02-阮氏水",
            ],
            lines,
        )
        self.assertEqual("MAIL-TEST-1_台中市政府勞工局_信封.docx", filename)

    def test_named_envelopes_fill_frontend_labels_and_drop_stale_text_box(self):
        record = {
            "mail_no": "MAIL-TEST-2",
            "recipient_postal_code_snapshot": "408221",
            "recipient_name_snapshot": "台中市政府勞工局",
            "recipient_department_snapshot": "入國通報",
            "recipient_contact_snapshot": "林小姐",
            "recipient_address_snapshot": "台中市南屯區精科路26號",
            "item_description": "D02-阮氏水",
            "mail_service": "限時掛號",
        }
        for template_name in ("信封(大).docx", "信封(小).docx"):
            with self.subTest(template_name=template_name):
                with patch.object(mail_model, "get_mail_record", return_value=record):
                    data, filename = mail_model.generate_envelope(1, template_name)

                document = Document(io.BytesIO(data))
                lines = [
                    paragraph.text.strip()
                    for paragraph in document.paragraphs
                    if paragraph.text.strip()
                ]
                expected = [
                    "408221 台中市南屯區精科路26號",
                    "(D02-阮氏水)",
                    "台中市政府勞工局 收",
                    "入國通報：林小姐",
                ]
                if template_name == "信封(大).docx":
                    expected.append("限時掛號")
                self.assertEqual(expected, lines)
                xml = document.element.xml
                self.assertNotIn("txbxContent", xml)
                self.assertNotIn("蔦松", xml)
                if template_name == "信封(小).docx":
                    self.assertNotIn("█", xml)
                self.assertEqual("MAIL-TEST-2_台中市政府勞工局_信封.docx", filename)

    def test_small_envelope_keeps_mark_only_for_registered_mail(self):
        base_record = {
            "mail_no": "MAIL-TEST-MARK",
            "recipient_postal_code_snapshot": "408221",
            "recipient_name_snapshot": "台中市政府勞工局",
            "recipient_address_snapshot": "台中市南屯區精科路26號",
            "item_description": "測試文件",
        }
        for mail_service, should_show_mark in (
            ("掛號", True),
            ("雙掛號", False),
            ("限時掛號", False),
            ("平信", False),
        ):
            with self.subTest(mail_service=mail_service):
                record = {**base_record, "mail_service": mail_service}
                with patch.object(mail_model, "get_mail_record", return_value=record):
                    data, _ = mail_model.generate_envelope(1, "信封(小).docx")

                output = Document(io.BytesIO(data))
                self.assertEqual(should_show_mark, "█" in output.element.xml)

    def test_recipient_department_contact_uses_colon_only_between_two_values(self):
        cases = (
            ({}, ""),
            ({"recipient_department_snapshot": "人事部"}, "人事部"),
            ({"recipient_contact_snapshot": "林小姐"}, "林小姐"),
            (
                {
                    "recipient_department_snapshot": "人事部",
                    "recipient_contact_snapshot": "林小姐",
                },
                "人事部：林小姐",
            ),
        )
        for record, expected in cases:
            with self.subTest(record=record):
                self.assertEqual(expected, mail_model._recipient_department_contact(record))

    def test_company_recipient_prints_phone_in_department_field(self):
        record = {
            "mail_no": "MAIL-COMPANY-RECIPIENT",
            "recipient_name_snapshot": "長青國際集團",
            "recipient_postal_code_snapshot": "505",
            "recipient_address_snapshot": "彰化縣鹿港鎮民權路163號3樓",
            "sender_type": "CUSTOM",
            "sender_name_snapshot": "客戶甲",
            "sender_address_snapshot": "彰化縣測試路1號",
            "item_description": "文件",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            document = Document()
            document.add_paragraph("{{recipient_department}}")
            document.save(template_dir / "代寄回本公司.docx")
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, _ = mail_model.generate_envelope(1, "代寄回本公司.docx")

        self.assertEqual("04-7785333", Document(io.BytesIO(data)).paragraphs[0].text)

    def test_company_recipient_prints_department_and_phone_together(self):
        record = {
            "recipient_name_snapshot": "長青國際集團",
            "recipient_address_snapshot": "彰化縣鹿港鎮民權路163號3樓",
            "recipient_department_snapshot": "人事部",
        }

        self.assertEqual(
            "人事部 04-7785333",
            mail_model._recipient_department_value(record),
        )

    def test_new_filename_automatically_uses_frontend_field_labels(self):
        record = {
            "mail_no": "MAIL-TEST-3",
            "recipient_postal_code_snapshot": "500",
            "recipient_name_snapshot": "測試公司",
            "recipient_department_snapshot": "人事部",
            "recipient_contact_snapshot": "陳先生",
            "recipient_address_snapshot": "彰化縣測試路1號",
            "item_description": "測試文件",
            "mail_service": "限時掛號",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            template = template_dir / "自訂橫式信封.docx"
            document = Document()
            document.add_paragraph("郵遞區號 收件地址")
            document.add_paragraph("(寄送內容)")
            document.add_paragraph("收件人／單位名稱 收")
            department_contact = document.add_paragraph()
            department_contact.add_run("收件")
            department_contact.add_run("部門：")
            department_contact.add_run("收件聯絡人")
            document.add_paragraph("郵件種類")
            document.save(template)
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, _ = mail_model.generate_envelope(1, template.name)

        output = Document(io.BytesIO(data))
        self.assertEqual(
            [
                "500 彰化縣測試路1號",
                "(測試文件)",
                "測試公司 收",
                "人事部：陳先生",
                "限時掛號",
            ],
            [paragraph.text for paragraph in output.paragraphs],
        )

    def test_frontend_field_labels_include_sender_details(self):
        record = {
            "mail_no": "MAIL-TEST-4",
            "sender_name_snapshot": "客戶甲",
            "sender_postal_code_snapshot": "104",
            "sender_address_snapshot": "台北市測試路2號",
            "sender_contact_snapshot": "王小姐",
            "sender_phone_snapshot": "02-1234-5678",
            "recipient_postal_code_snapshot": "500",
            "recipient_name_snapshot": "測試公司",
            "recipient_address_snapshot": "彰化縣測試路1號",
            "mail_service": "掛號",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            template = template_dir / "含寄件人.docx"
            document = Document()
            document.add_paragraph("寄件人名稱 寄件人郵遞區號 寄件人地址")
            document.add_paragraph("寄件聯絡人 寄件人電話")
            document.add_paragraph("郵遞區號 收件人／單位名稱 收件地址")
            document.save(template)
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, _ = mail_model.generate_envelope(1, template.name)

        output = Document(io.BytesIO(data))
        self.assertEqual(
            [
                "客戶甲 104 台北市測試路2號",
                "王小姐 02-1234-5678",
                "500 測試公司 彰化縣測試路1號",
            ],
            [paragraph.text for paragraph in output.paragraphs],
        )

    def test_sender_postal_code_accepts_postal_alias(self):
        record = {
            "mail_no": "MAIL-TEST-5",
            "sender_postal_code_snapshot": "104",
            "recipient_name_snapshot": "測試公司",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            template = template_dir / "寄件人郵寄區號別名.docx"
            document = Document()
            document.add_paragraph("寄件人郵寄區號")
            document.save(template)
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, _ = mail_model.generate_envelope(1, template.name)

        output = Document(io.BytesIO(data))
        self.assertEqual("104", output.paragraphs[0].text)

    def test_envelope_template_replaces_paragraph_and_table_placeholders(self):
        record = {
            "mail_no": "MAIL-TEST-1",
            "recipient_postal_code_snapshot": "500",
            "recipient_name_snapshot": "測試公司",
            "recipient_department_snapshot": "人事部",
            "recipient_address_snapshot": "彰化縣測試路1號",
            "sender_name_snapshot": "王小明",
            "item_description": "測試文件",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            template = template_dir / "信封.docx"
            document = Document()
            document.add_paragraph("{{recipient_name}} {{recipient_address}}")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "{{item_description}}"
            document.save(template)
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, filename = mail_model.generate_envelope(1, "信封.docx")

        output = Document(io.BytesIO(data))
        self.assertIn("測試公司", output.paragraphs[0].text)
        self.assertIn("彰化縣測試路1號", output.paragraphs[0].text)
        self.assertEqual("測試文件", output.tables[0].cell(0, 0).text)
        self.assertEqual("MAIL-TEST-1_測試公司_信封.docx", filename)

    def test_multiple_envelopes_are_combined_one_per_page_in_input_order(self):
        records = {
            1: {
                "mail_no": "MAIL-TEST-1",
                "recipient_name_snapshot": "第一家公司",
                "recipient_address_snapshot": "第一個地址",
                "item_description": "第一份文件",
                "mail_service": "掛號",
            },
            2: {
                "mail_no": "MAIL-TEST-2",
                "recipient_name_snapshot": "第二家公司",
                "recipient_address_snapshot": "第二個地址",
                "item_description": "第二份文件",
                "mail_service": "限時掛號",
            },
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            template = template_dir / "批次信封.docx"
            document = Document()
            document.add_paragraph("{{recipient_name}}｜{{recipient_address}}")
            document.add_paragraph("{{item_description}}｜{{mail_service}}")
            document.save(template)
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", side_effect=records.get),
            ):
                data, filename = mail_model.generate_envelopes([2, 1], template.name)

        output = Document(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in output.paragraphs)
        self.assertLess(text.index("第二家公司"), text.index("第一家公司"))
        self.assertIn("第二份文件｜限時掛號", text)
        self.assertIn("第一份文件｜掛號", text)
        self.assertEqual(0, len(output.element.body.xpath('.//w:br[@w:type="page"]')))
        self.assertEqual(1, len(output.element.body.xpath(".//w:pageBreakBefore")))
        self.assertEqual("信封套印_2件.docx", filename)

    def test_multiple_envelopes_require_at_least_one_record(self):
        with self.assertRaisesRegex(ValueError, "至少選擇一筆寄件"):
            mail_model.generate_envelopes([], "B4.docx")

    def test_history_excel_accepts_postgresql_timezone_datetime(self):
        records = [
            {
                "id": 1,
                "mail_no": "MAIL-TEST-1",
                "created_at": datetime(2026, 10, 14, 1, 30, tzinfo=timezone.utc),
                "operator_name_snapshot": "陳小姐",
                "processor_name": "王小明",
                "sender_type": "CLIENT",
                "sender_name_snapshot": "客戶甲",
                "recipient_name_snapshot": "測試公司",
                "recipient_department_snapshot": "人事部",
                "recipient_contact_snapshot": "李小姐",
                "recipient_address_snapshot": "測試地址",
                "item_description": "測試文件",
                "mail_service": "掛號",
                "billing_employer_names": "政泰、台彰",
                "status": "READY",
                "batch_no": None,
                "tracking_no": None,
                "shipping_fee": None,
                "notes": "請附回郵信封",
            }
        ]
        dataframe = mail_view._records_dataframe(records)
        self.assertEqual("王小明", dataframe.loc[0, "處理人"])
        self.assertEqual("人事部", dataframe.loc[0, "收件部門"])
        self.assertEqual("李小姐", dataframe.loc[0, "收件聯絡人"])
        self.assertEqual("政泰、台彰", dataframe.loc[0, "關聯客戶"])
        self.assertEqual("請附回郵信封", dataframe.loc[0, "備註"])
        self.assertEqual("備註", dataframe.columns[-1])
        self.assertEqual("陳小姐", dataframe.loc[0, "寄件人"])
        self.assertEqual("代寄：客戶甲", dataframe.loc[0, "寄件名義"])
        self.assertIsNone(dataframe.loc[0, "建立時間"].tzinfo)
        self.assertEqual(9, dataframe.loc[0, "建立時間"].hour)
        output = io.BytesIO()
        dataframe.to_excel(output, index=False)
        self.assertGreater(len(output.getvalue()), 0)

        my_records_dataframe = mail_view._records_dataframe(records, include_identifiers=False)
        self.assertNotIn("id", my_records_dataframe.columns)
        self.assertNotIn("寄件編號", my_records_dataframe.columns)

    def test_user_mail_records_supports_keyword_and_status_filters(self):
        with patch.object(mail_model, "_fetch_all", return_value=[]) as fetch_all:
            mail_model.get_user_mail_records(7, "測試關鍵字", "DISPATCHED")

        query, params = fetch_all.call_args.args
        self.assertIn("r.created_by_user_id = %s", query)
        self.assertNotIn("r.mail_no ILIKE %s", query)
        self.assertIn("b.batch_no ILIKE %s", query)
        self.assertIn("bi.tracking_no ILIKE %s", query)
        self.assertIn("r.recipient_name_snapshot ILIKE %s", query)
        self.assertIn("r.sender_name_snapshot ILIKE %s", query)
        self.assertIn("r.item_description ILIKE %s", query)
        self.assertIn("mail_record_billing_employers", query)
        self.assertIn("be.employer_id_snapshot", query)
        self.assertIn("be.employer_name_snapshot", query)
        self.assertIn("r.status = %s", query)
        self.assertEqual(
            (7, *(["%測試關鍵字%"] * 6), "DISPATCHED", 100),
            params,
        )

    def test_user_mail_records_keeps_owner_filter_without_search_filters(self):
        with patch.object(mail_model, "_fetch_all", return_value=[]) as fetch_all:
            mail_model.get_user_mail_records(9)

        query, params = fetch_all.call_args.args
        self.assertIn("WHERE r.created_by_user_id = %s", query)
        self.assertNotIn("ILIKE", query)
        self.assertEqual((9, 100), params)

    def test_all_records_can_search_and_return_operator_name(self):
        with patch.object(mail_model, "_fetch_all", return_value=[]) as fetch_all:
            mail_model.search_all_records("陳小姐", "DISPATCHED")

        query, params = fetch_all.call_args.args
        self.assertIn("r.operator_name_snapshot ILIKE %s", query)
        self.assertIn("r.operator_name_snapshot", query.split("FROM mail_records", 1)[0])
        self.assertEqual(8, params.count("%陳小姐%"))
        self.assertEqual("DISPATCHED", params[-2])


    def test_company_sender_stores_no_sender_fields(self):
        sender, error = mail_model._normalize_sender({"sender_name": "應被忽略"})
        self.assertIsNone(error)
        self.assertEqual("COMPANY", sender["sender_type"])
        self.assertIsNone(sender["sender_name"])
        self.assertIsNone(sender["sender_postal_code"])
        self.assertIsNone(sender["sender_address"])

    def test_on_behalf_sender_requires_name_and_address(self):
        _, error = mail_model._normalize_sender({"sender_type": "CUSTOM", "sender_name": "客戶甲"})
        self.assertIn("必填", error)
        _, error = mail_model._normalize_sender({"sender_type": "BOGUS"})
        self.assertIn("不正確", error)

    def test_custom_sender_drops_workplace_link(self):
        sender, error = mail_model._normalize_sender(
            {
                "sender_type": "CUSTOM",
                "sender_employer_workplace_id": 9,
                "sender_name": " 客戶甲 ",
                "sender_address": "台中市測試路1號",
                "sender_phone": "04-1234567",
            }
        )
        self.assertIsNone(error)
        self.assertIsNone(sender["sender_employer_workplace_id"])
        self.assertEqual("客戶甲", sender["sender_name"])
        self.assertEqual("04-1234567", sender["sender_phone"])

    def test_create_record_saves_operator_and_client_sender_separately(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.side_effect = [None, {"id": 5}]
        data = {
            "recipient_type": "CUSTOM",
            "recipient_name": "勞工局",
            "recipient_address": "台中市南屯區精科路26號",
            "item_description": "聘僱許可函",
            "sender_type": "CLIENT",
            "sender_employer_workplace_id": 7,
            "sender_name": "測試公司",
            "sender_postal_code": "506",
            "sender_address": "彰化縣福興鄉福工路2之1號",
            "sender_contact": "林小姐",
            "sender_phone": "04-7777777",
            "planned_shipping_fee": 36,
        }
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message, record_id = mail_model.create_mail_record(1, "王小明", data)

        self.assertTrue(success, message)
        self.assertEqual(5, record_id)
        insert_sql, params = cursor.execute.call_args_list[1].args
        self.assertEqual(insert_sql.count("%s"), len(params))
        self.assertEqual(
            (1, "王小明", "CLIENT", 7, "測試公司", "506", "彰化縣福興鄉福工路2之1號", "林小姐", "04-7777777"),
            params[:9],
        )
        self.assertEqual(36, params[-2])

    def test_create_record_rejects_decimal_postage_before_database_access(self):
        data = {
            "recipient_type": "CUSTOM",
            "recipient_name": "勞工局",
            "recipient_address": "台中市南屯區精科路26號",
            "item_description": "聘僱許可函",
            "planned_shipping_fee": "28.5",
        }
        with patch.object(mail_model.database, "get_db_connection") as get_connection:
            success, message, record_id = mail_model.create_mail_record(1, "王小明", data)

        self.assertFalse(success)
        self.assertIn("不接受小數", message)
        self.assertIsNone(record_id)
        get_connection.assert_not_called()

    def test_create_record_saves_multiple_billing_employers(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.side_effect = [None, {"id": 8}]
        data = {
            "recipient_type": "CUSTOM",
            "recipient_name": "安心會計師事務所",
            "recipient_address": "彰化市中山路1號",
            "item_description": "稅務文件",
            "billing_employers": [
                {
                    "key": "CODE:a001",
                    "employer_workplace_id": 11,
                    "employer_id": "A001",
                    "employer_name": "東方公司",
                },
                {
                    "key": "CODE:a018",
                    "employer_workplace_id": 18,
                    "employer_id": "A018",
                    "employer_name": "永盛公司",
                },
            ],
        }
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message, record_id = mail_model.create_mail_record(1, "王小明", data)

        self.assertTrue(success, message)
        self.assertEqual(8, record_id)
        mail_insert = next(
            call for call in cursor.execute.call_args_list
            if "INSERT INTO mail_records" in call.args[0]
        )
        self.assertEqual((11, "A001", "東方公司"), mail_insert.args[1][17:20])
        relation_inserts = [
            call.args[1]
            for call in cursor.execute.call_args_list
            if "INSERT INTO mail_record_billing_employers" in call.args[0]
        ]
        self.assertEqual(
            [
                (8, "CODE:a001", 11, "A001", "東方公司", 1),
                (8, "CODE:a018", 18, "A018", "永盛公司", 2),
            ],
            relation_inserts,
        )

    def test_create_mail_records_creates_all_selected_recipients_in_one_transaction(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.side_effect = [None, None, {"id": 21}, {"id": 22}]
        rows = [
            {
                "recipient_type": "CUSTOM",
                "recipient_name": "第一家公司",
                "recipient_address": "第一個地址",
                "item_description": "年度文件",
                "mail_service": "掛號",
            },
            {
                "recipient_type": "CUSTOM",
                "recipient_name": "第二家公司",
                "recipient_address": "第二個地址",
                "item_description": "年度文件",
                "mail_service": "掛號",
            },
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message, record_ids = mail_model.create_mail_records(1, "王小明", rows)

        self.assertTrue(success, message)
        self.assertEqual([21, 22], record_ids)
        inserts = [
            call for call in cursor.execute.call_args_list
            if "INSERT INTO mail_records" in call.args[0]
        ]
        self.assertEqual(2, len(inserts))
        connection.commit.assert_called_once_with()
        connection.rollback.assert_not_called()

    def test_multi_mail_editor_creates_one_payload_per_non_empty_row(self):
        billing_option = {
            "key": "CODE:a001",
            "employer_workplace_id": 11,
            "employer_id": "A001",
            "employer_name": "東方公司",
            "label": "A001｜東方公司",
        }
        rows = [
            {
                "寄件名義": "本公司",
                "收件人／單位名稱": "勞動部",
                "收件地址": "臺北市某路 1 號",
                "寄送內容": "甲員聘僱文件",
                "郵件種類": "掛號",
                "關聯客戶": ["A001｜東方公司"],
            },
            {
                "寄件名義": "代客戶",
                "寄件人名稱": "東方公司",
                "寄件人地址": "彰化縣某路 2 號",
                "收件人／單位名稱": "勞動部",
                "收件地址": "臺北市某路 1 號",
                "寄送內容": "乙員聘僱文件",
                "郵件種類": "限時掛號",
                "關聯客戶": [],
            },
            {"寄件名義": "本公司", "郵件種類": "掛號", "關聯客戶": []},
        ]

        payloads = mail_view._multi_mail_rows_to_payload(
            rows, {billing_option["label"]: billing_option}
        )

        self.assertEqual(2, len(payloads))
        self.assertEqual(
            ["甲員聘僱文件", "乙員聘僱文件"],
            [row["item_description"] for row in payloads],
        )
        self.assertEqual("COMPANY", payloads[0]["sender_type"])
        self.assertEqual([billing_option], payloads[0]["billing_employers"])
        self.assertEqual("CUSTOM", payloads[1]["sender_type"])
        self.assertEqual("東方公司", payloads[1]["sender_name"])

    def test_multi_mail_editor_can_fill_sender_and_recipient_from_options(self):
        billing_option = {
            "key": "CODE:a001",
            "employer_workplace_id": 11,
            "employer_id": "A001",
            "employer_name": "東方公司",
            "label": "A001｜東方公司",
        }
        sender_option = {
            "type": "CLIENT",
            "source_id": 11,
            "name": "東方公司",
            "postal_code": "500",
            "address": "彰化縣某路 2 號",
            "contact": "王小姐",
            "phone": "04-1234567",
            "employer_id": "A001",
        }
        recipient_option = {
            "type": "GOVERNMENT",
            "source_id": 9,
            "name": "勞動部",
            "postal_code": "100",
            "address": "臺北市某路 1 號",
            "department": "收文科",
            "contact": "李先生",
            "related_employer_keys": ["CODE:a001"],
        }
        rows = [
            {
                # Streamlit may wrap selectbox cells in a one-item list.
                "帶入寄件人": ["A001｜東方公司｜彰化縣某路 2 號"],
                "帶入收件人": ["勞動部｜臺北市某路 1 號"],
                "收件聯絡人": "陳小姐",
                "寄送內容": "聘僱文件",
                "郵件種類": "掛號",
                "關聯客戶": [["A001｜東方公司"]],
            }
        ]

        payloads = mail_view._multi_mail_rows_to_payload(
            rows,
            {billing_option["label"]: billing_option},
            {"A001｜東方公司｜彰化縣某路 2 號": sender_option},
            {"勞動部｜臺北市某路 1 號": recipient_option},
        )

        self.assertEqual(1, len(payloads))
        payload = payloads[0]
        self.assertEqual("CLIENT", payload["sender_type"])
        self.assertEqual(11, payload["sender_employer_workplace_id"])
        self.assertEqual("東方公司", payload["sender_name"])
        self.assertEqual("勞動部", payload["recipient_name"])
        self.assertEqual("陳小姐", payload["recipient_contact"])
        self.assertEqual(9, payload["government_agency_id"])
        self.assertEqual([billing_option], payload["billing_employers"])

    def test_multi_mail_editor_displays_selected_sender_and_recipient_fields(self):
        sender_option = {
            "type": "CLIENT",
            "name": "東方公司",
            "postal_code": "500",
            "address": "彰化縣某路 2 號",
            "contact": "王小姐",
            "phone": "04-1234567",
        }
        recipient_option = {
            "type": "GOVERNMENT",
            "name": "勞動部",
            "postal_code": "100",
            "address": "臺北市某路 1 號",
            "department": "收文科",
            "contact": "李先生",
        }
        rows = [
            {
                "帶入寄件人": ["東方公司選項"],
                "帶入收件人": ["勞動部選項"],
                "寄送內容": "聘僱文件",
                "_已帶入寄件人": "本公司",
                "_已帶入收件人": None,
            }
        ]

        hydrated, changed = mail_view._hydrate_multi_mail_editor_rows(
            rows,
            {"東方公司選項": sender_option},
            {"勞動部選項": recipient_option},
        )

        self.assertTrue(changed)
        self.assertEqual("東方公司", hydrated[0]["寄件人名稱"])
        self.assertEqual("彰化縣某路 2 號", hydrated[0]["寄件人地址"])
        self.assertEqual("勞動部", hydrated[0]["收件人／單位名稱"])
        self.assertEqual("收文科", hydrated[0]["收件部門"])
        self.assertEqual("臺北市某路 1 號", hydrated[0]["收件地址"])

        unchanged, changed_again = mail_view._hydrate_multi_mail_editor_rows(
            hydrated,
            {"東方公司選項": sender_option},
            {"勞動部選項": recipient_option},
        )
        self.assertFalse(changed_again)
        self.assertEqual(hydrated, unchanged)

    def test_multi_mail_editor_defaults_new_rows_to_company_sender(self):
        hydrated, changed = mail_view._hydrate_multi_mail_editor_rows(
            [{"寄送內容": "稅務文件"}],
            {"本公司": {"type": "COMPANY"}},
            {},
        )

        self.assertTrue(changed)
        self.assertEqual("本公司", hydrated[0]["帶入寄件人"])
        self.assertEqual("本公司", hydrated[0]["_已帶入寄件人"])
        self.assertEqual("", hydrated[0]["寄件人名稱"])
        self.assertEqual("", hydrated[0]["寄件人郵遞區號"])
        self.assertEqual("", hydrated[0]["寄件人地址"])

    def test_create_mail_records_keeps_every_row_as_an_independent_record(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.side_effect = [None, None, {"id": 31}, {"id": 32}]
        rows = [
            {"recipient_name": "同一單位", "recipient_address": "相同地址", "item_description": "文件甲"},
            {"recipient_name": "同一單位", "recipient_address": "相同地址", "item_description": "文件乙"},
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message, record_ids = mail_model.create_mail_records(
                1, "王小明", rows, source="MULTI_ENTRY"
            )

        self.assertTrue(success, message)
        self.assertEqual([31, 32], record_ids)
        inserts = [
            call for call in cursor.execute.call_args_list
            if "INSERT INTO mail_records" in call.args[0]
        ]
        self.assertEqual(2, len(inserts))
        self.assertEqual(2, audit.call_count)
        self.assertTrue(
            all(call.kwargs["after"]["source"] == "MULTI_ENTRY" for call in audit.call_args_list)
        )

    def test_create_mail_records_rolls_back_all_when_one_recipient_is_duplicate(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.side_effect = [None, {"mail_no": "MAIL-OLD"}]
        rows = [
            {"recipient_name": "甲", "recipient_address": "地址甲", "item_description": "文件"},
            {"recipient_name": "乙", "recipient_address": "地址乙", "item_description": "文件"},
        ]
        with patch.object(mail_model.database, "get_db_connection", return_value=connection):
            success, message, record_ids = mail_model.create_mail_records(1, "王小明", rows)

        self.assertFalse(success)
        self.assertIn("MAIL-OLD", message)
        self.assertEqual([], record_ids)
        self.assertFalse(any("INSERT INTO mail_records" in call.args[0] for call in cursor.execute.call_args_list))
        connection.rollback.assert_called_once_with()

    def test_create_batch_preserves_dragged_order_and_removes_duplicates(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [{"id": 3}, {"id": 1}, {"id": 2}]
        cursor.fetchone.return_value = {"id": 99}
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message, batch_id = mail_model.create_batch(
                [3, 1, 3, 2], "REGULAR", 1, posting_date=date(2026, 10, 14)
            )

        self.assertTrue(success, message)
        self.assertEqual(99, batch_id)
        item_insert_params = [
            call.args[1]
            for call in cursor.execute.call_args_list
            if "INSERT INTO mail_batch_items" in call.args[0]
        ]
        self.assertEqual(
            [(99, 3, 1, None), (99, 1, 2, None), (99, 2, 3, None)],
            item_insert_params,
        )

    def test_waiting_batch_can_change_from_bulk_to_regular(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {
            "id": 99,
            "batch_no": "POST-20261014-00099",
            "status": "CREATED",
            "submission_mode": "BULK_FORM",
        }
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.update_batch_submission_mode(99, "REGULAR", 7)

        self.assertTrue(success, message)
        update_call = next(
            call for call in cursor.execute.call_args_list
            if "SET submission_mode" in call.args[0]
        )
        self.assertEqual(("REGULAR", None, None, 99), update_call.args[1])
        self.assertEqual(
            {"submission_mode": "BULK_FORM"}, audit.call_args.kwargs["before"]
        )
        self.assertEqual(
            {"submission_mode": "REGULAR"}, audit.call_args.kwargs["after"]
        )
        connection.commit.assert_called_once_with()

    def test_waiting_batch_can_change_from_regular_to_bulk(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {
            "id": 99,
            "batch_no": "POST-20261014-00099",
            "status": "CREATED",
            "submission_mode": "REGULAR",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            template = template_dir / "郵局大宗掛號單.xlsx"
            template.write_bytes(b"test template")
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model.database, "get_db_connection", return_value=connection),
                patch.object(mail_model, "_audit"),
            ):
                success, message = mail_model.update_batch_submission_mode(99, "BULK_FORM", 7)

        self.assertTrue(success, message)
        update_call = next(
            call for call in cursor.execute.call_args_list
            if "SET submission_mode" in call.args[0]
        )
        mode, filename, version, batch_id = update_call.args[1]
        self.assertEqual(("BULK_FORM", "郵局大宗掛號單.xlsx", 99), (mode, filename, batch_id))
        self.assertEqual(16, len(version))
        connection.commit.assert_called_once_with()

    def test_dispatched_batch_cannot_change_submission_mode(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {
            "id": 99,
            "batch_no": "POST-20261014-00099",
            "status": "DISPATCHED",
            "submission_mode": "REGULAR",
        }
        with patch.object(mail_model.database, "get_db_connection", return_value=connection):
            success, message = mail_model.update_batch_submission_mode(99, "BULK_FORM", 7)

        self.assertFalse(success)
        self.assertIn("等待結案", message)
        self.assertEqual(1, cursor.execute.call_count)
        connection.commit.assert_not_called()

    def test_void_batch_returns_items_and_preserves_batch_postage(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {"status": "CREATED"}
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message = mail_model.void_batch(99, 7, "重新併批")

        self.assertTrue(success, message)
        return_sql = next(
            call.args[0]
            for call in cursor.execute.call_args_list
            if "UPDATE mail_records r" in call.args[0]
        )
        self.assertIn("planned_shipping_fee = COALESCE(bi.shipping_fee", return_sql)
        connection.commit.assert_called_once_with()

    def test_envelope_template_fills_sender_placeholders(self):
        record = {
            "mail_no": "MAIL-TEST-2",
            "recipient_name_snapshot": "勞工局",
            "recipient_address_snapshot": "台中市南屯區精科路26號",
            "sender_type": "CLIENT",
            "sender_name_snapshot": "測試公司",
            "sender_postal_code_snapshot": "506",
            "sender_address_snapshot": "彰化縣福興鄉福工路2之1號",
            "sender_contact_snapshot": "林小姐",
            "sender_phone_snapshot": "04-7777777",
            "item_description": "文件",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            document = Document()
            document.add_paragraph("{{sender_postal_code}} {{sender_address}}")
            document.add_paragraph("{{sender_name}} {{sender_contact}} {{sender_phone}}")
            document.save(template_dir / "代寄.docx")
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, _ = mail_model.generate_envelope(1, "代寄.docx")

        output = Document(io.BytesIO(data))
        self.assertEqual("506 彰化縣福興鄉福工路2之1號", output.paragraphs[0].text)
        self.assertEqual("測試公司 林小姐 04-7777777", output.paragraphs[1].text)

    def test_company_sender_fills_sender_placeholders(self):
        record = {
            "mail_no": "MAIL-TEST-3",
            "recipient_name_snapshot": "勞工局",
            "recipient_address_snapshot": "地址",
            "sender_type": "COMPANY",
            "item_description": "文件",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            template_dir = Path(temp_dir)
            document = Document()
            document.add_paragraph(
                "[{{sender_postal_code}}{{sender_name}}{{sender_address}}{{sender_phone}}]"
            )
            document.save(template_dir / "代寄.docx")
            with (
                patch.object(mail_model, "TEMPLATE_DIR", template_dir),
                patch.object(mail_model, "get_mail_record", return_value=record),
            ):
                data, _ = mail_model.generate_envelope(1, "代寄.docx")

        self.assertEqual(
            "[505長青國際集團彰化縣鹿港鎮民權路163號3樓04-7785333]",
            Document(io.BytesIO(data)).paragraphs[0].text,
        )

    def test_bulk_workbook_prints_backfilled_posting_date(self):
        batch = {
            "batch_no": "POST-20261001-00001",
            "submission_mode": "BULK_FORM",
            "template_filename": "郵局大宗掛號單.xlsx",
            "posting_date": date(2026, 10, 1),
        }
        with patch.object(mail_model, "get_batch", return_value=(batch, self._items(1))):
            data, _ = mail_model.generate_bulk_workbook(1)
        sheet = load_workbook(io.BytesIO(data))["P1"]
        self.assertEqual((115, 10, 1), (sheet["C2"].value, sheet["E2"].value, sheet["G2"].value))

    def test_create_batch_rejects_future_posting_date(self):
        future = date.today() + timedelta(days=1)
        success, message, batch_id = mail_model.create_batch([1], "REGULAR", 1, posting_date=future)
        self.assertFalse(success)
        self.assertIn("不可晚於今天", message)
        self.assertIsNone(batch_id)

    def test_receipt_file_named_by_batch_no_without_overwriting(self):
        upload = MagicMock()
        upload.name = "收據.PDF"
        upload.getvalue.return_value = b"%PDF-1.4 test"
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.object(mail_model, "RECEIPT_DIR", Path(temp_dir)):
                first = mail_model.save_receipt(upload, "POST-20261001-00001")
                second = mail_model.save_receipt(upload, "POST-20261001-00001")
        self.assertEqual("POST-20261001-00001.pdf", Path(first["path"]).name)
        self.assertEqual("POST-20261001-00001-2.pdf", Path(second["path"]).name)
        self.assertEqual("收據.PDF", first["original_name"])

    def _edit_data(self):
        return {
            "recipient_name": "勞工局",
            "recipient_address": "台中市南屯區精科路26號",
            "item_description": "聘僱許可函（更正）",
            "planned_shipping_fee": 36,
        }

    def test_update_record_allowed_before_dispatch(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        before = {
            "id": 5, "mail_no": "MAIL-1", "created_by_user_id": 1, "status": "BATCHED",
            "item_description": "聘僱許可函",
        }
        cursor.fetchone.side_effect = [
            before,
            {**before, "item_description": "聘僱許可函（更正）", "planned_shipping_fee": 36},
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.update_mail_record(5, 1, self._edit_data())

        self.assertTrue(success, message)
        update_sql, params = cursor.execute.call_args_list[1].args
        self.assertEqual(update_sql.count("%s"), len(params))
        self.assertEqual(36, params[-4])
        connection.commit.assert_called_once_with()
        self.assertEqual(
            {"item_description": "聘僱許可函（更正）", "planned_shipping_fee": 36},
            audit.call_args.kwargs["after"],
        )

    def test_update_record_blocked_after_dispatch_or_by_other_user(self):
        for status, owner, expected in (("DISPATCHED", 1, "不可再修改"), ("READY", 2, "自己建立")):
            connection = MagicMock()
            cursor = connection.cursor.return_value.__enter__.return_value
            cursor.fetchone.return_value = {
                "id": 5, "mail_no": "MAIL-1", "created_by_user_id": owner, "status": status,
            }
            with patch.object(mail_model.database, "get_db_connection", return_value=connection):
                success, message = mail_model.update_mail_record(5, 1, self._edit_data())
            self.assertFalse(success)
            self.assertIn(expected, message)
            self.assertEqual(1, cursor.execute.call_count)
            connection.commit.assert_not_called()

    def test_cancel_mail_record_allows_postal_operator_to_cancel_other_owners_ready_mail(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {"mail_no": "MAIL-1"}
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.cancel_mail_record(
                5, 9, "重複寄件", allow_any_owner=True
            )

        self.assertTrue(success, message)
        query, params = cursor.execute.call_args.args
        self.assertNotIn("created_by_user_id", query)
        self.assertIn("status IN ('DRAFT', 'READY')", query)
        self.assertEqual((9, "重複寄件", 5), params)
        connection.commit.assert_called_once_with()
        audit.assert_called_once_with(
            cursor,
            "MAIL_RECORD",
            5,
            "CANCEL",
            9,
            after={"reason": "重複寄件"},
        )

    def test_cancel_mail_record_still_limits_regular_user_to_own_mail(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = None
        with patch.object(mail_model.database, "get_db_connection", return_value=connection):
            success, message = mail_model.cancel_mail_record(5, 9, "")

        self.assertFalse(success)
        query, params = cursor.execute.call_args.args
        self.assertIn("created_by_user_id = %s", query)
        self.assertEqual((9, "使用者取消", 5, 9), params)
        connection.commit.assert_not_called()

    def test_restore_cancelled_mail_record_returns_it_to_ready(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {"mail_no": "MAIL-1"}
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.restore_cancelled_mail_record(5, 1)

        self.assertTrue(success, message)
        query, params = cursor.execute.call_args.args
        self.assertIn("status = 'READY'", query)
        self.assertIn("cancelled_at = NULL", query)
        self.assertIn("cancelled_by_user_id = NULL", query)
        self.assertIn("cancel_reason = NULL", query)
        self.assertIn("status = 'CANCELLED'", query)
        self.assertEqual((5, 1), params)
        connection.commit.assert_called_once_with()
        audit.assert_called_once_with(
            cursor,
            "MAIL_RECORD",
            5,
            "RESTORE_READY",
            1,
            before={"status": "CANCELLED"},
            after={"status": "READY"},
        )

    def test_restore_cancelled_mail_record_rejects_non_cancelled_or_other_owner(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = None
        with patch.object(mail_model.database, "get_db_connection", return_value=connection):
            success, message = mail_model.restore_cancelled_mail_record(5, 1)

        self.assertFalse(success)
        self.assertIn("只能恢復自己已取消", message)
        connection.commit.assert_not_called()

    def test_update_dispatched_batch_changes_history_without_requiring_new_receipt(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {
            "id": 7,
            "status": "DISPATCHED",
            "total_shipping_fee": 72,
            "posting_date": date(2026, 10, 1),
            "receipt_file_path": "mail_receipts/original.pdf",
        }
        updates = [
            {"batch_item_id": 21, "tracking_no": " 0012345 ", "shipping_fee": 38}
        ]
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit") as audit,
        ):
            success, message = mail_model.update_dispatched_batch(
                7, 1, 80, updates, date(2026, 10, 2)
            )

        self.assertTrue(success, message)
        item_sql, item_params = cursor.execute.call_args_list[1].args
        self.assertIn("UPDATE mail_batch_items", item_sql)
        self.assertEqual(("0012345", 38, 21, 7), item_params)
        batch_sql, batch_params = cursor.execute.call_args_list[2].args
        self.assertNotIn("receipt_file_path", batch_sql)
        self.assertEqual((80, date(2026, 10, 2), 7), batch_params)
        self.assertEqual(
            "mail_receipts/original.pdf",
            audit.call_args.kwargs["after"]["receipt_file_path"],
        )
        connection.commit.assert_called_once_with()

    def test_update_dispatched_batch_can_replace_receipt(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {
            "id": 7,
            "status": "DISPATCHED",
            "total_shipping_fee": 72,
            "posting_date": date(2026, 10, 1),
            "receipt_file_path": "mail_receipts/original.pdf",
        }
        receipt = {
            "path": "mail_receipts/replacement.pdf",
            "original_name": "更正收據.pdf",
            "content_type": "application/pdf",
        }
        with (
            patch.object(mail_model.database, "get_db_connection", return_value=connection),
            patch.object(mail_model, "_audit"),
        ):
            success, message = mail_model.update_dispatched_batch(
                7, 1, 80, [], date(2026, 10, 2), receipt
            )

        self.assertTrue(success, message)
        batch_sql, batch_params = cursor.execute.call_args_list[1].args
        self.assertIn("receipt_file_path", batch_sql)
        self.assertEqual(
            (
                80,
                date(2026, 10, 2),
                "mail_receipts/replacement.pdf",
                "更正收據.pdf",
                "application/pdf",
                1,
                7,
            ),
            batch_params,
        )

    def test_update_dispatched_batch_rejects_unfinished_batch(self):
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = {"id": 7, "status": "CREATED"}
        with patch.object(mail_model.database, "get_db_connection", return_value=connection):
            success, message = mail_model.update_dispatched_batch(
                7, 1, 80, [], date(2026, 10, 2)
            )

        self.assertFalse(success)
        self.assertIn("只有已結案", message)
        connection.commit.assert_not_called()


if __name__ == "__main__":
    unittest.main()
