import base64
from copy import deepcopy
import hashlib
import hmac
import io
import math
import os
import re
import secrets
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from psycopg2.extras import Json

import database
from postal_code_lookup import infer_postal_code


BASE_PATH = Path(database.BASE_PATH)
TEMPLATE_DIR = BASE_PATH / "doc_templates"
RECEIPT_DIR = BASE_PATH / "mail_receipts"
ADMIN_CODE_KEY = "admin_code_hash"
MAX_RECEIPT_BYTES = 15 * 1024 * 1024
ALLOWED_RECEIPTS = {
    ".pdf": "application/pdf",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
}
COMPANY_SENDER = {
    "type": "COMPANY",
    "name": "長青國際集團",
    "postal_code": "505",
    "address": "彰化縣鹿港鎮民權路163號3樓",
    "contact": "",
    "phone": "04-7785333",
}


def _clean(value):
    if isinstance(value, float) and math.isnan(value):
        return None
    text = " ".join(str(value or "").split()).strip()
    return text or None


def _parse_shipping_fee(value):
    """Normalize a single Streamlit/Pandas value into an integer TWD amount."""
    if value is None:
        return None
    while isinstance(value, (list, tuple)):
        if not value:
            return None
        if len(value) != 1:
            raise ValueError(f"無法辨識郵資「{value}」，請輸入單一數字。")
        value = value[0]
    if hasattr(value, "item"):
        try:
            value = value.item()
        except (ValueError, TypeError):
            pass
    if isinstance(value, float) and math.isnan(value):
        return None
    text = str(value).strip()
    if not text or text.casefold() in {"nan", "none", "<na>"}:
        return None
    bracketed = re.fullmatch(r"\[\s*([^,\[\]]+)\s*\]", text)
    if bracketed:
        text = bracketed.group(1).strip()
    text = (
        text.replace("NT$", "")
        .replace("nt$", "")
        .replace("$", "")
        .replace(",", "")
        .replace("，", "")
        .strip()
    )
    if not text:
        return None
    try:
        fee = Decimal(text)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"無法辨識郵資「{value}」，請輸入數字。")
    if not fee.is_finite():
        return None
    if fee < 0:
        raise ValueError("郵資不可小於 0。")
    if fee != fee.to_integral_value():
        raise ValueError("台幣郵資不接受小數，請輸入整數。")
    return int(fee)


def _fetch_all(query, params=()):
    conn = database.get_db_connection()
    if not conn:
        return []
    try:
        with conn.cursor() as cursor:
            cursor.execute(query, params)
            return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()


IMPORTANT_AUDITS = frozenset(
    {
        ("MAIL_RECORD", "UPDATE"),
        ("MAIL_RECORD", "CANCEL"),
        ("MAIL_RECORD", "RESTORE_READY"),
        ("MAIL_BATCH", "CREATE"),
        ("MAIL_BATCH", "CHANGE_SUBMISSION_MODE"),
        ("MAIL_BATCH", "DISPATCH"),
        ("MAIL_BATCH", "UPDATE_DISPATCHED"),
        ("MAIL_BATCH", "VOID"),
        ("MAIL_SETTINGS", "SET_ADMIN_CODE"),
    }
)


def _audit(cursor, entity_type, entity_id, action, actor_user_id, before=None, after=None):
    if (entity_type, action) not in IMPORTANT_AUDITS:
        return
    cursor.execute(
        """
        INSERT INTO mail_audit_logs
            (entity_type, entity_id, action, actor_user_id, before_data, after_data)
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (
            entity_type,
            entity_id,
            action,
            actor_user_id,
            Json(before) if before is not None else None,
            Json(after) if after is not None else None,
        ),
    )


# ---------- 操作人員與共通管理碼 ----------

def get_mail_users(include_inactive=False):
    where = "" if include_inactive else "WHERE is_active = TRUE"
    return _fetch_all(
        f"SELECT id, display_name, is_active, sort_order FROM mail_users {where} "
        "ORDER BY sort_order, display_name"
    )


def add_mail_user(display_name, actor_user_id=None):
    name = _clean(display_name)
    if not name:
        return False, "請輸入人員姓名。"
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO mail_users (display_name)
                VALUES (%s)
                ON CONFLICT (display_name)
                DO UPDATE SET is_active = TRUE, updated_at = NOW()
                RETURNING id
                """,
                (name,),
            )
            user_id = cursor.fetchone()["id"]
            if actor_user_id:
                _audit(cursor, "MAIL_USER", user_id, "UPSERT", actor_user_id, after={"display_name": name})
        conn.commit()
        return True, "人員名單已更新。"
    except Exception as exc:
        conn.rollback()
        return False, f"新增人員失敗：{exc}"
    finally:
        conn.close()


def set_mail_user_active(user_id, is_active, actor_user_id):
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "UPDATE mail_users SET is_active = %s, updated_at = NOW() WHERE id = %s",
                (bool(is_active), user_id),
            )
            _audit(cursor, "MAIL_USER", user_id, "SET_ACTIVE", actor_user_id, after={"is_active": bool(is_active)})
        conn.commit()
        return True, "人員狀態已更新。"
    except Exception as exc:
        conn.rollback()
        return False, f"更新人員失敗：{exc}"
    finally:
        conn.close()


def save_mail_users(rows, actor_user_id):
    normalized = []
    duplicate_names = set()
    for index, row in enumerate(rows, 1):
        user_id = row.get("id")
        if isinstance(user_id, float) and math.isnan(user_id):
            user_id = None
        name = _clean(row.get("姓名"))
        if user_id is None and not name:
            continue
        if not name:
            return False, f"第 {index} 列的人員姓名不可空白。"

        name_key = name.casefold()
        if name_key in duplicate_names:
            return False, f"第 {index} 列與其他列重複：{name}。"
        duplicate_names.add(name_key)

        raw_active = row.get("啟用", True)
        if raw_active is None or (isinstance(raw_active, float) and math.isnan(raw_active)):
            is_active = True
        else:
            is_active = bool(raw_active)
        raw_sort_order = row.get("排序", 0)
        if raw_sort_order is None or (isinstance(raw_sort_order, float) and math.isnan(raw_sort_order)):
            sort_order = 0
        else:
            try:
                sort_order = int(raw_sort_order)
            except (TypeError, ValueError):
                return False, f"第 {index} 列的排序必須是整數。"
        normalized.append(
            {
                "id": int(user_id) if user_id is not None else None,
                "display_name": name,
                "is_active": is_active,
                "sort_order": sort_order,
            }
        )

    if not normalized or not any(row["is_active"] for row in normalized):
        return False, "至少需要保留一位啟用中的操作人員。"

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT id FROM mail_users FOR UPDATE")
            existing_ids = {row["id"] for row in cursor.fetchall()}
            submitted_ids = {row["id"] for row in normalized if row["id"] is not None}
            unknown_ids = submitted_ids - existing_ids
            if unknown_ids:
                raise ValueError("資料已被其他人變更，請重新整理後再試。")

            inserted = updated = 0
            for row in normalized:
                if row["id"] is None:
                    cursor.execute(
                        """
                        INSERT INTO mail_users (display_name, is_active, sort_order)
                        VALUES (%s, %s, %s)
                        """,
                        (row["display_name"], row["is_active"], row["sort_order"]),
                    )
                    inserted += 1
                else:
                    cursor.execute(
                        """
                        UPDATE mail_users SET
                            display_name = %s,
                            is_active = %s,
                            sort_order = %s,
                            updated_at = NOW()
                        WHERE id = %s
                        """,
                        (
                            row["display_name"],
                            row["is_active"],
                            row["sort_order"],
                            row["id"],
                        ),
                    )
                    updated += 1

            removed_ids = existing_ids - submitted_ids
            if removed_ids:
                cursor.execute(
                    """
                    UPDATE mail_users
                    SET is_active = FALSE, updated_at = NOW()
                    WHERE id = ANY(%s)
                    """,
                    (list(removed_ids),),
                )
            _audit(
                cursor,
                "MAIL_USER",
                0,
                "BATCH_SAVE",
                actor_user_id,
                after={
                    "inserted": inserted,
                    "updated": updated,
                    "deactivated_by_removal": len(removed_ids),
                },
            )
        conn.commit()
        return True, f"操作人員已儲存：新增 {inserted} 筆、更新 {updated} 筆、停用 {len(removed_ids)} 筆。"
    except Exception as exc:
        conn.rollback()
        return False, f"儲存操作人員失敗：{exc}"
    finally:
        conn.close()


def _hash_admin_code(code, salt=None, iterations=260000):
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", code.encode("utf-8"), salt, iterations)
    return "pbkdf2_sha256${}${}${}".format(
        iterations,
        base64.urlsafe_b64encode(salt).decode("ascii"),
        base64.urlsafe_b64encode(digest).decode("ascii"),
    )


def has_admin_code():
    rows = _fetch_all("SELECT 1 FROM mail_settings WHERE setting_key = %s", (ADMIN_CODE_KEY,))
    return bool(rows)


def verify_admin_code(code):
    rows = _fetch_all(
        "SELECT setting_value FROM mail_settings WHERE setting_key = %s",
        (ADMIN_CODE_KEY,),
    )
    if not rows or not code:
        return False
    try:
        algorithm, iterations, salt_text, digest_text = rows[0]["setting_value"].split("$", 3)
        if algorithm != "pbkdf2_sha256":
            return False
        candidate = _hash_admin_code(
            code,
            base64.urlsafe_b64decode(salt_text.encode("ascii")),
            int(iterations),
        )
        return hmac.compare_digest(candidate, rows[0]["setting_value"])
    except (ValueError, TypeError):
        return False


def set_admin_code(new_code, actor_user_id, current_code=None):
    if not new_code or len(new_code) < 4:
        return False, "管理碼至少需要 4 個字元。"
    if has_admin_code() and not verify_admin_code(current_code):
        return False, "目前管理碼不正確。"
    encoded = _hash_admin_code(new_code)
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO mail_settings (setting_key, setting_value, updated_at, updated_by_user_id)
                VALUES (%s, %s, NOW(), %s)
                ON CONFLICT (setting_key) DO UPDATE SET
                    setting_value = EXCLUDED.setting_value,
                    updated_at = NOW(),
                    updated_by_user_id = EXCLUDED.updated_by_user_id
                """,
                (ADMIN_CODE_KEY, encoded, actor_user_id),
            )
            _audit(cursor, "MAIL_SETTINGS", 0, "SET_ADMIN_CODE", actor_user_id)
        conn.commit()
        return True, "共通管理碼已設定。"
    except Exception as exc:
        conn.rollback()
        return False, f"設定管理碼失敗：{exc}"
    finally:
        conn.close()


def log_admin_event(user_id, action):
    conn = database.get_db_connection()
    if not conn:
        return
    try:
        with conn.cursor() as cursor:
            _audit(cursor, "ADMIN_SESSION", user_id, action, user_id)
        conn.commit()
    except Exception:
        conn.rollback()
    finally:
        conn.close()


# ---------- 收件人與寄件 ----------

def _employer_relation_key(employer_id, employer_name):
    employer_id = _clean(employer_id)
    employer_name = _clean(employer_name)
    if employer_id:
        return f"CODE:{employer_id.casefold()}"
    return f"NAME:{(employer_name or '').casefold()}"


def get_employer_relation_options():
    """取得供機構關聯使用的雇主清單；同一代號的多個工作地址只顯示一次。"""
    rows = _fetch_all(
        """
        SELECT id, employer_id, employer_name
        FROM "EmployerWorkplaces"
        WHERE employer_name IS NOT NULL AND BTRIM(employer_name) <> ''
        ORDER BY employer_id NULLS LAST, employer_name
        """
    )
    options = []
    seen = set()
    for row in rows:
        employer_id = _clean(row.get("employer_id"))
        employer_name = _clean(row.get("employer_name"))
        key = _employer_relation_key(employer_id, employer_name)
        if not employer_name or key in seen:
            continue
        seen.add(key)
        options.append(
            {
                "key": key,
                "employer_workplace_id": row.get("id"),
                "employer_id": employer_id,
                "employer_name": employer_name,
                "label": f"{employer_id}｜{employer_name}" if employer_id else employer_name,
            }
        )
    return options


def get_company_recipient_option():
    """Fixed recipient shortcut shown only for mail sent on a client's behalf."""
    return {
        "key": "COMPANY",
        "type": "CUSTOM",
        "source_id": None,
        "label": f"本公司｜{COMPANY_SENDER['name']}｜{COMPANY_SENDER['address']}",
        "name": COMPANY_SENDER["name"],
        "address": COMPANY_SENDER["address"],
        "postal_code": COMPANY_SENDER["postal_code"],
        "department": "",
        "contact": "",
        "related_employer_keys": [],
        "employer_id": None,
    }


def get_recipient_options(user_id=None):
    employers = _fetch_all(
        """
        SELECT id, employer_id, employer_name, address, postal_code
        FROM "EmployerWorkplaces"
        WHERE address IS NOT NULL AND BTRIM(address) <> ''
        ORDER BY employer_id NULLS LAST, employer_name, address
        """
    )
    agencies = _fetch_all(
        """
        SELECT a.id, a.agency_name, a.department_name, a.contact_person,
               a.postal_code, a.address,
               COALESCE(
                   (
                       SELECT ARRAY_AGG(
                           CASE
                               WHEN ae.employer_id IS NOT NULL AND BTRIM(ae.employer_id) <> ''
                               THEN ae.employer_id || '｜' || ae.employer_name
                               ELSE ae.employer_name
                           END
                           ORDER BY ae.employer_id NULLS LAST, ae.employer_name
                       )
                       FROM mail_government_agency_employers ae
                       WHERE ae.agency_id = a.id
                   ),
                   ARRAY[]::TEXT[]
               ) AS related_employers,
               COALESCE(
                   (
                       SELECT ARRAY_AGG(ae.employer_key ORDER BY ae.employer_id NULLS LAST, ae.employer_name)
                       FROM mail_government_agency_employers ae
                       WHERE ae.agency_id = a.id
                   ),
                   ARRAY[]::TEXT[]
               ) AS related_employer_keys
        FROM mail_government_agencies a
        WHERE a.is_active = TRUE
        ORDER BY a.agency_name, a.department_name
        """
    )
    options = []
    # 機關優先顯示，其次才是雇主。
    for row in agencies:
        department = f"｜{row['department_name']}" if row.get("department_name") else ""
        contact = _clean(row.get("contact_person"))
        related_employers = row.get("related_employers") or []
        label_parts = [f"{row['agency_name']}{department}", row["address"]]
        if contact:
            label_parts.append(contact)
        if related_employers:
            label_parts.append(f"關聯：{'、'.join(related_employers)}")
        options.append(
            {
                "key": f"G:{row['id']}",
                "type": "GOVERNMENT",
                "source_id": row["id"],
                "label": "｜".join(label_parts),
                "name": row["agency_name"],
                "address": row["address"],
                "postal_code": row.get("postal_code") or "",
                "department": row.get("department_name") or "",
                "contact": contact or "",
                "related_employer_keys": row.get("related_employer_keys") or [],
                "employer_id": None,
            }
        )
    for row in employers:
        code = row.get("employer_id") or "無編號"
        options.append(
            {
                "key": f"E:{row['id']}",
                "type": "EMPLOYER",
                "source_id": row["id"],
                "label": f"{code}｜{row['employer_name']}｜{row['address']}",
                "name": row["employer_name"],
                "address": row["address"],
                "postal_code": row.get("postal_code") or "",
                "department": "",
                "contact": "",
                "employer_key": _employer_relation_key(row.get("employer_id"), row["employer_name"]),
                "related_employer_keys": [],
                "employer_id": row.get("employer_id"),
            }
        )
    if user_id is not None:
        personal_recipients = _fetch_all(
            """
            SELECT i.id, i.recipient_type, i.employer_workplace_id,
                   i.government_agency_id, i.recipient_name,
                   i.recipient_department, i.recipient_postal_code,
                   i.recipient_address, i.recipient_contact,
                   i.billing_employers, l.list_name
            FROM mail_recipient_list_items i
            JOIN mail_recipient_lists l ON l.id = i.list_id
            WHERE l.owner_user_id = %s
            ORDER BY l.list_name, i.sequence_no
            """,
            (user_id,),
        )
        personal_options = []
        for row in personal_recipients:
            billing_employers = row.get("billing_employers") or []
            billing_keys = [
                employer.get("key")
                for employer in billing_employers
                if employer.get("key")
            ]
            recipient_type = row.get("recipient_type") or "CUSTOM"
            if recipient_type == "EMPLOYER":
                source_id = row.get("employer_workplace_id")
            elif recipient_type == "GOVERNMENT":
                source_id = row.get("government_agency_id")
            else:
                source_id = None
            personal_options.append(
                {
                    "key": f"L:{row['id']}",
                    "type": recipient_type,
                    "source_id": source_id,
                    "label": (
                        f"我的名單｜{row['list_name']}｜{row['recipient_name']}｜"
                        f"{row['recipient_address']}"
                    ),
                    "name": row["recipient_name"],
                    "address": row["recipient_address"],
                    "postal_code": row.get("recipient_postal_code") or "",
                    "department": row.get("recipient_department") or "",
                    "contact": row.get("recipient_contact") or "",
                    "employer_key": billing_keys[0] if recipient_type == "EMPLOYER" and billing_keys else None,
                    "related_employer_keys": billing_keys,
                    "employer_id": (
                        billing_employers[0].get("employer_id")
                        if billing_employers else None
                    ),
                }
            )
        options = personal_options + options
    return options


# ---------- 我的寄送名單 ----------

RECIPIENT_LIST_COLUMNS = ("recipient_name", "recipient_postal_code", "recipient_address", "recipient_department", "recipient_contact", "notes")


def _normalize_recipient_list_items(items):
    normalized = []
    seen = set()
    for index, item in enumerate(items or [], 1):
        name = _clean(item.get("recipient_name"))
        address = _clean(item.get("recipient_address"))
        optional_values = [item.get(field) for field in RECIPIENT_LIST_COLUMNS if field not in {"recipient_name", "recipient_address"}]
        if not name and not address and not any(_clean(value) for value in optional_values):
            continue
        if not name or not address:
            return None, f"第 {index} 筆的收件人名稱與地址不可空白。"
        recipient_type = item.get("recipient_type") or "CUSTOM"
        if recipient_type not in {"EMPLOYER", "GOVERNMENT", "CUSTOM"}:
            recipient_type = "CUSTOM"
        duplicate_key = (name.casefold(), address.casefold())
        if duplicate_key in seen:
            return None, f"第 {index} 筆與名單內其他收件人重複：{name}。"
        seen.add(duplicate_key)
        billing_employers, error = _normalize_billing_employers(
            {"billing_employers": item.get("billing_employers") or []}
        )
        if error:
            return None, f"第 {index} 筆：{error}"
        normalized.append(
            {
                "recipient_type": recipient_type,
                "employer_workplace_id": item.get("employer_workplace_id"),
                "government_agency_id": item.get("government_agency_id"),
                "recipient_name": name,
                "recipient_department": _clean(item.get("recipient_department")),
                "recipient_postal_code": _clean(item.get("recipient_postal_code")),
                "recipient_address": address,
                "recipient_contact": _clean(item.get("recipient_contact")),
                "billing_employers": billing_employers,
                "notes": _clean(item.get("notes")),
            }
        )
    if not normalized:
        return None, "名單至少需要一位收件人。"
    return normalized, None


def recipient_option_to_list_item(option, employer_relations=None):
    """Convert an existing recipient option into an editable list snapshot."""
    employer_relations = employer_relations or {}
    billing_keys = []
    if option.get("type") == "EMPLOYER" and option.get("employer_key"):
        billing_keys = [option["employer_key"]]
    elif option.get("type") == "GOVERNMENT":
        billing_keys = option.get("related_employer_keys") or []
    return {
        "recipient_type": option.get("type") or "CUSTOM",
        "employer_workplace_id": option.get("source_id") if option.get("type") == "EMPLOYER" else None,
        "government_agency_id": option.get("source_id") if option.get("type") == "GOVERNMENT" else None,
        "recipient_name": option.get("name") or "",
        "recipient_department": option.get("department") or "",
        "recipient_postal_code": option.get("postal_code") or "",
        "recipient_address": option.get("address") or "",
        "recipient_contact": option.get("contact") or "",
        "billing_employers": [employer_relations[key] for key in billing_keys if key in employer_relations],
        "notes": "",
    }


def parse_recipient_list_excel(data):
    """Read the first worksheet of an xlsx recipient list into canonical rows."""
    if not data:
        raise ValueError("Excel 檔案是空的。")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"無法讀取 Excel 檔案：{exc}") from exc
    try:
        sheet = workbook.worksheets[0]
        raw_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not raw_rows:
        raise ValueError("Excel 沒有資料。")

    aliases = {
        "recipient_name": {"收件人名稱", "收件人", "名稱", "單位名稱", "收件單位"},
        "recipient_postal_code": {"郵遞區號", "郵編", "郵遞區號（選填）"},
        "recipient_address": {"地址", "收件地址"},
        "recipient_department": {"部門", "收件部門"},
        "recipient_contact": {"聯絡人", "收件聯絡人"},
        "notes": {"備註"},
    }
    normalized_headers = ["".join(str(value or "").split()) for value in raw_rows[0]]
    field_indexes = {}
    for field, names in aliases.items():
        for index, header in enumerate(normalized_headers):
            if header in {"".join(name.split()) for name in names}:
                field_indexes[field] = index
                break
    if "recipient_name" not in field_indexes or "recipient_address" not in field_indexes:
        raise ValueError("Excel 必須包含「收件人名稱」與「地址」欄位。")

    items = []
    for values in raw_rows[1:]:
        item = {"recipient_type": "CUSTOM", "billing_employers": []}
        for field in RECIPIENT_LIST_COLUMNS:
            index = field_indexes.get(field)
            value = values[index] if index is not None and index < len(values) else None
            if field == "recipient_postal_code" and isinstance(value, float) and value.is_integer():
                value = str(int(value))
            item[field] = _clean(value) or ""
        if any(item[field] for field in RECIPIENT_LIST_COLUMNS):
            items.append(item)
    normalized, error = _normalize_recipient_list_items(items)
    if error:
        raise ValueError(error)
    return normalized


def add_recipient_items_to_agencies(items, actor_user_id):
    """Add imported recipients to the shared agency directory, skipping duplicates."""
    normalized, error = _normalize_recipient_list_items(items)
    if error:
        return False, error
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, agency_name, department_name, address
                FROM mail_government_agencies
                FOR UPDATE
                """
            )
            existing_keys = {
                (
                    (_clean(row.get("agency_name")) or "").casefold(),
                    (_clean(row.get("department_name")) or "").casefold(),
                    (_clean(row.get("address")) or "").casefold(),
                )
                for row in cursor.fetchall()
            }
            inserted = skipped = 0
            for item in normalized:
                key = (
                    item["recipient_name"].casefold(),
                    (item["recipient_department"] or "").casefold(),
                    item["recipient_address"].casefold(),
                )
                if key in existing_keys:
                    skipped += 1
                    continue
                cursor.execute(
                    """
                    INSERT INTO mail_government_agencies (
                        agency_name, department_name, contact_person,
                        postal_code, address, is_active
                    ) VALUES (%s, %s, %s, %s, %s, TRUE)
                    """,
                    (
                        item["recipient_name"], item["recipient_department"],
                        item["recipient_contact"], item["recipient_postal_code"],
                        item["recipient_address"],
                    ),
                )
                existing_keys.add(key)
                inserted += 1
            _audit(
                cursor,
                "MAIL_AGENCY",
                0,
                "IMPORT_FROM_RECIPIENT_LIST",
                actor_user_id,
                after={"inserted": inserted, "skipped": skipped},
            )
        conn.commit()
        return True, f"常用收件機構：新增 {inserted} 筆、略過既有 {skipped} 筆。"
    except Exception as exc:
        conn.rollback()
        return False, f"新增常用收件機構失敗：{exc}"
    finally:
        conn.close()


def get_recipient_lists(user_id):
    return _fetch_all(
        """
        SELECT l.id, l.list_name, l.created_at, l.updated_at, COUNT(i.id)::INTEGER AS item_count
        FROM mail_recipient_lists l
        LEFT JOIN mail_recipient_list_items i ON i.list_id = l.id
        WHERE l.owner_user_id = %s
        GROUP BY l.id
        ORDER BY l.updated_at DESC, l.list_name
        """,
        (user_id,),
    )


def get_recipient_list(list_id, user_id):
    headers = _fetch_all(
        "SELECT * FROM mail_recipient_lists WHERE id = %s AND owner_user_id = %s",
        (list_id, user_id),
    )
    if not headers:
        return None, []
    items = _fetch_all(
        """
        SELECT id, sequence_no, recipient_type, employer_workplace_id,
               government_agency_id, recipient_name, recipient_department,
               recipient_postal_code, recipient_address, recipient_contact,
               billing_employers, notes
        FROM mail_recipient_list_items
        WHERE list_id = %s
        ORDER BY sequence_no
        """,
        (list_id,),
    )
    return headers[0], items


def save_recipient_list(user_id, list_name, items, list_id=None):
    list_name = _clean(list_name)
    if not list_name:
        return False, "請輸入名單名稱。", None
    normalized, error = _normalize_recipient_list_items(items)
    if error:
        return False, error, None
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。", None
    try:
        with conn.cursor() as cursor:
            if list_id is None:
                cursor.execute(
                    """
                    INSERT INTO mail_recipient_lists (list_name, owner_user_id)
                    VALUES (%s, %s) RETURNING id
                    """,
                    (list_name, user_id),
                )
                list_id = cursor.fetchone()["id"]
                action = "CREATE"
            else:
                cursor.execute(
                    """
                    UPDATE mail_recipient_lists
                    SET list_name = %s, updated_at = NOW()
                    WHERE id = %s AND owner_user_id = %s
                    RETURNING id
                    """,
                    (list_name, list_id, user_id),
                )
                if not cursor.fetchone():
                    raise ValueError("找不到名單，或您沒有修改權限。")
                cursor.execute("DELETE FROM mail_recipient_list_items WHERE list_id = %s", (list_id,))
                action = "UPDATE"
            for sequence_no, item in enumerate(normalized, 1):
                cursor.execute(
                    """
                    INSERT INTO mail_recipient_list_items (
                        list_id, sequence_no, recipient_type, employer_workplace_id,
                        government_agency_id, recipient_name, recipient_department,
                        recipient_postal_code, recipient_address, recipient_contact,
                        billing_employers, notes
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        list_id, sequence_no, item["recipient_type"], item["employer_workplace_id"],
                        item["government_agency_id"], item["recipient_name"], item["recipient_department"],
                        item["recipient_postal_code"], item["recipient_address"], item["recipient_contact"],
                        Json(item["billing_employers"]), item["notes"],
                    ),
                )
            _audit(cursor, "MAIL_RECIPIENT_LIST", list_id, action, user_id, after={"name": list_name, "count": len(normalized)})
        conn.commit()
        return True, f"名單「{list_name}」已儲存，共 {len(normalized)} 位收件人。", list_id
    except Exception as exc:
        conn.rollback()
        return False, f"儲存名單失敗：{exc}", None
    finally:
        conn.close()


def add_government_agency(name, address, department, postal_code, actor_user_id):
    name = _clean(name)
    address = _clean(address)
    if not name or not address:
        return False, "機關名稱與地址為必填。"
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO mail_government_agencies
                    (agency_name, department_name, postal_code, address)
                VALUES (%s, %s, %s, %s) RETURNING id
                """,
                (name, _clean(department), _clean(postal_code), address),
            )
            agency_id = cursor.fetchone()["id"]
            _audit(cursor, "MAIL_AGENCY", agency_id, "CREATE", actor_user_id, after={"name": name})
        conn.commit()
        return True, "機關已加入收件人名單。"
    except Exception as exc:
        conn.rollback()
        return False, f"新增機關失敗：{exc}"
    finally:
        conn.close()


def get_government_agencies_for_editor():
    return _fetch_all(
        """
        SELECT a.id, a.is_active, a.agency_name, a.department_name,
               a.contact_person, a.postal_code, a.address,
               COALESCE(
                   (
                       SELECT ARRAY_AGG(
                           CASE
                               WHEN ae.employer_id IS NOT NULL AND BTRIM(ae.employer_id) <> ''
                               THEN ae.employer_id || '｜' || ae.employer_name
                               ELSE ae.employer_name
                           END
                           ORDER BY ae.employer_id NULLS LAST, ae.employer_name
                       )
                       FROM mail_government_agency_employers ae
                       WHERE ae.agency_id = a.id
                   ),
                   ARRAY[]::TEXT[]
               ) AS related_employers
        FROM mail_government_agencies a
        ORDER BY a.is_active DESC, a.agency_name, a.department_name, a.address
        """
    )


def save_government_agencies(rows, actor_user_id):
    normalized = []
    duplicate_keys = set()
    for index, row in enumerate(rows, 1):
        agency_id = row.get("id")
        if isinstance(agency_id, float) and math.isnan(agency_id):
            agency_id = None
        name = _clean(row.get("機關名稱"))
        department = _clean(row.get("部門"))
        contact_person = _clean(row.get("聯絡人"))
        postal_code = _clean(row.get("郵遞區號"))
        address = _clean(row.get("地址"))
        raw_relations = row.get("關聯雇主") if "關聯雇主" in row else None
        if isinstance(raw_relations, float) and math.isnan(raw_relations):
            raw_relations = []
        if raw_relations is not None and not isinstance(raw_relations, (list, tuple, set)):
            return False, f"第 {index} 列的關聯雇主格式不正確。"
        raw_active = row.get("啟用", True)
        if raw_active is None or (isinstance(raw_active, float) and math.isnan(raw_active)):
            is_active = True
        else:
            is_active = bool(raw_active)
        if not any((name, department, contact_person, postal_code, address, raw_relations)):
            continue
        if not name or not address:
            return False, f"第 {index} 列的機關名稱與地址不可空白。"
        key = (name.casefold(), (department or "").casefold(), address.casefold())
        if key in duplicate_keys:
            return False, f"第 {index} 列與其他列重複：{name}。"
        duplicate_keys.add(key)
        normalized.append(
            {
                "id": int(agency_id) if agency_id is not None else None,
                "is_active": is_active,
                "agency_name": name,
                "department_name": department,
                "contact_person": contact_person,
                "postal_code": postal_code,
                "address": address,
                "related_employer_labels": list(raw_relations) if raw_relations is not None else None,
            }
        )

    relation_options = None
    if any(row["related_employer_labels"] is not None for row in normalized):
        relation_options = {
            option["label"]: option for option in get_employer_relation_options()
        }
        for index, row in enumerate(normalized, 1):
            labels = row["related_employer_labels"]
            if labels is None:
                row["related_employers"] = None
                continue
            unknown = [label for label in labels if label not in relation_options]
            if unknown:
                return False, f"第 {index} 列含有不存在的關聯雇主：{unknown[0]}。請重新整理後再試。"
            row["related_employers"] = [relation_options[label] for label in labels]

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, is_active, agency_name, department_name, contact_person,
                       postal_code, address
                FROM mail_government_agencies
                FOR UPDATE
                """
            )
            existing_rows = {
                row["id"]: {
                    "is_active": bool(row["is_active"]),
                    "agency_name": _clean(row["agency_name"]),
                    "department_name": _clean(row["department_name"]),
                    "contact_person": _clean(row.get("contact_person")),
                    "postal_code": _clean(row["postal_code"]),
                    "address": _clean(row["address"]),
                }
                for row in cursor.fetchall()
            }
            existing_ids = set(existing_rows)
            submitted_ids = {row["id"] for row in normalized if row["id"] is not None}
            unknown_ids = submitted_ids - existing_ids
            if unknown_ids:
                raise ValueError("資料已被其他人變更，請重新整理後再試。")

            existing_relations = {}
            if relation_options is not None:
                cursor.execute(
                    """
                    SELECT agency_id, employer_key, employer_id, employer_name
                    FROM mail_government_agency_employers
                    WHERE agency_id = ANY(%s)
                    """,
                    (list(existing_ids),),
                )
                for relation in cursor.fetchall():
                    existing_relations.setdefault(relation["agency_id"], []).append(
                        {
                            "key": relation["employer_key"],
                            "employer_id": _clean(relation.get("employer_id")),
                            "employer_name": _clean(relation.get("employer_name")),
                        }
                    )

            inserted = updated = 0
            for row in normalized:
                if row["id"] is None:
                    cursor.execute(
                        """
                        INSERT INTO mail_government_agencies
                            (agency_name, department_name, contact_person, postal_code, address, is_active)
                        VALUES (%s, %s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        (
                            row["agency_name"],
                            row["department_name"],
                            row["contact_person"],
                            row["postal_code"],
                            row["address"],
                            row["is_active"],
                        ),
                    )
                    row["id"] = cursor.fetchone()["id"]
                    inserted += 1
                else:
                    before = existing_rows[row["id"]]
                    base_changed = not all(
                        before[field] == row[field]
                        for field in (
                            "is_active",
                            "agency_name",
                            "department_name",
                            "contact_person",
                            "postal_code",
                            "address",
                        )
                    )
                    if base_changed:
                        cursor.execute(
                            """
                            UPDATE mail_government_agencies SET
                                agency_name = %s,
                                department_name = %s,
                                contact_person = %s,
                                postal_code = %s,
                                address = %s,
                                is_active = %s,
                                updated_at = NOW()
                            WHERE id = %s
                            """,
                            (
                                row["agency_name"],
                                row["department_name"],
                                row["contact_person"],
                                row["postal_code"],
                                row["address"],
                                row["is_active"],
                                row["id"],
                            ),
                        )

                    desired_relations = row.get("related_employers")
                    relation_changed = False
                    if desired_relations is not None:
                        current_keys = {
                            relation["key"] for relation in existing_relations.get(row["id"], [])
                        }
                        desired_keys = {relation["key"] for relation in desired_relations}
                        relation_changed = current_keys != desired_keys
                    if base_changed or relation_changed:
                        updated += 1

                desired_relations = row.get("related_employers")
                if desired_relations is not None:
                    current_keys = {
                        relation["key"] for relation in existing_relations.get(row["id"], [])
                    }
                    desired_keys = {relation["key"] for relation in desired_relations}
                    if current_keys != desired_keys:
                        cursor.execute(
                            "DELETE FROM mail_government_agency_employers WHERE agency_id = %s",
                            (row["id"],),
                        )
                        for relation in desired_relations:
                            cursor.execute(
                                """
                                INSERT INTO mail_government_agency_employers
                                    (agency_id, employer_key, employer_id, employer_name)
                                VALUES (%s, %s, %s, %s)
                                """,
                                (
                                    row["id"],
                                    relation["key"],
                                    relation["employer_id"],
                                    relation["employer_name"],
                                ),
                            )

            removed_ids = {
                agency_id
                for agency_id in existing_ids - submitted_ids
                if existing_rows[agency_id]["is_active"]
            }
            if removed_ids:
                cursor.execute(
                    """
                    UPDATE mail_government_agencies
                    SET is_active = FALSE, updated_at = NOW()
                    WHERE id = ANY(%s)
                    """,
                    (list(removed_ids),),
                )
            _audit(
                cursor,
                "MAIL_AGENCY",
                0,
                "BATCH_SAVE",
                actor_user_id,
                after={
                    "inserted": inserted,
                    "updated": updated,
                    "deactivated_by_removal": len(removed_ids),
                },
            )
        conn.commit()
        return True, f"機關已儲存：新增 {inserted} 筆、更新 {updated} 筆、停用 {len(removed_ids)} 筆。"
    except Exception as exc:
        conn.rollback()
        return False, f"儲存機關失敗：{exc}"
    finally:
        conn.close()


# ---------- 實際寄件人（代客寄件） ----------

SENDER_TYPES = {"COMPANY", "CLIENT", "CUSTOM"}


def get_sender_options(recent_days=365):
    """客戶寄件人清單：雇主工作地點，以及近期用過的自訂寄件人。

    聯絡人與電話沿用該寄件人最近一次寄件時填寫的內容。
    """
    employers = _fetch_all(
        """
        SELECT w.id, w.employer_id, w.employer_name, w.address, w.postal_code,
               last.sender_contact_snapshot AS contact,
               last.sender_phone_snapshot AS phone
        FROM "EmployerWorkplaces" w
        LEFT JOIN LATERAL (
            SELECT r.sender_contact_snapshot, r.sender_phone_snapshot
            FROM mail_records r
            WHERE r.sender_employer_workplace_id = w.id
            ORDER BY r.created_at DESC LIMIT 1
        ) last ON TRUE
        WHERE w.address IS NOT NULL AND BTRIM(w.address) <> ''
        ORDER BY w.employer_id NULLS LAST, w.employer_name, w.address
        """
    )
    recent_custom = _fetch_all(
        """
        SELECT DISTINCT ON (sender_name_snapshot, sender_address_snapshot)
               sender_name_snapshot AS name, sender_address_snapshot AS address,
               sender_postal_code_snapshot AS postal_code,
               sender_contact_snapshot AS contact, sender_phone_snapshot AS phone
        FROM mail_records
        WHERE sender_type = 'CUSTOM'
          AND created_at >= NOW() - (%s * INTERVAL '1 day')
        ORDER BY sender_name_snapshot, sender_address_snapshot, created_at DESC
        """,
        (recent_days,),
    )
    options = []
    for row in employers:
        code = row.get("employer_id") or "無編號"
        options.append(
            {
                "key": f"E:{row['id']}",
                "type": "CLIENT",
                "source_id": row["id"],
                "label": f"{code}｜{row['employer_name']}｜{row['address']}",
                "name": row["employer_name"],
                "postal_code": row.get("postal_code") or "",
                "address": row["address"],
                "contact": row.get("contact") or "",
                "phone": row.get("phone") or "",
                "employer_id": row.get("employer_id"),
            }
        )
    for index, row in enumerate(recent_custom):
        options.append(
            {
                "key": f"C:{index}:{row['name']}|{row['address']}",
                "type": "CUSTOM",
                "source_id": None,
                "label": f"（曾用）{row['name']}｜{row['address']}",
                "name": row["name"],
                "postal_code": row.get("postal_code") or "",
                "address": row["address"],
                "contact": row.get("contact") or "",
                "phone": row.get("phone") or "",
                "employer_id": None,
            }
        )
    return options


def _normalize_sender(data):
    """回傳 (sender_dict, error)。本公司寄件不保存寄件人欄位。"""
    sender_type = data.get("sender_type") or "COMPANY"
    if sender_type not in SENDER_TYPES:
        return None, "寄件人類型不正確。"
    if sender_type == "COMPANY":
        return {
            "sender_type": "COMPANY",
            "sender_employer_workplace_id": None,
            "sender_name": None,
            "sender_postal_code": None,
            "sender_address": None,
            "sender_contact": None,
            "sender_phone": None,
        }, None
    sender = {
        "sender_type": sender_type,
        "sender_employer_workplace_id": (
            data.get("sender_employer_workplace_id") if sender_type == "CLIENT" else None
        ),
        "sender_name": _clean(data.get("sender_name")),
        "sender_postal_code": _clean(data.get("sender_postal_code")),
        "sender_address": _clean(data.get("sender_address")),
        "sender_contact": _clean(data.get("sender_contact")),
        "sender_phone": _clean(data.get("sender_phone")),
    }
    if not sender["sender_name"] or not sender["sender_address"]:
        return None, "代客寄件時，寄件人名稱與地址為必填。"
    return sender, None


def _normalize_billing_employers(data):
    """Normalize the multi-select value while accepting the legacy single-employer fields."""
    raw_employers = data.get("billing_employers")
    if raw_employers is None:
        legacy_name = _clean(data.get("billing_employer_name"))
        raw_employers = (
            [
                {
                    "employer_workplace_id": data.get("billing_employer_workplace_id"),
                    "employer_id": data.get("billing_employer_id"),
                    "employer_name": legacy_name,
                }
            ]
            if legacy_name
            else []
        )
    if not isinstance(raw_employers, (list, tuple)):
        return None, "關聯客戶格式不正確。"

    normalized = []
    seen = set()
    for employer in raw_employers:
        if not isinstance(employer, dict):
            return None, "關聯客戶格式不正確。"
        employer_id = _clean(employer.get("employer_id"))
        employer_name = _clean(employer.get("employer_name") or employer.get("name"))
        if not employer_name:
            return None, "關聯客戶名稱不可空白。"
        key = _employer_relation_key(employer_id, employer_name)
        if key in seen:
            continue
        seen.add(key)
        normalized.append(
            {
                "key": key,
                "employer_workplace_id": employer.get("employer_workplace_id") or employer.get("source_id"),
                "employer_id": employer_id,
                "employer_name": employer_name,
            }
        )
    return normalized, None


def _replace_billing_employers(cursor, record_id, employers, delete_existing=True):
    if delete_existing:
        cursor.execute(
            "DELETE FROM mail_record_billing_employers WHERE mail_record_id = %s",
            (record_id,),
        )
    for sequence_no, employer in enumerate(employers, 1):
        cursor.execute(
            """
            INSERT INTO mail_record_billing_employers (
                mail_record_id, employer_key, employer_workplace_id,
                employer_id_snapshot, employer_name_snapshot, sequence_no
            ) VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                record_id,
                employer["key"],
                employer["employer_workplace_id"],
                employer["employer_id"],
                employer["employer_name"],
                sequence_no,
            ),
        )


def create_mail_record(user_id, operator_name, data):
    name = _clean(data.get("recipient_name"))
    address = _clean(data.get("recipient_address"))
    description = _clean(data.get("item_description"))
    if not name or not address or not description:
        return False, "收件人、地址與寄送內容皆為必填。", None
    sender, error = _normalize_sender(data)
    if error:
        return False, error, None
    try:
        planned_shipping_fee = _parse_shipping_fee(data.get("planned_shipping_fee"))
    except ValueError as exc:
        return False, str(exc), None
    billing_employers, error = _normalize_billing_employers(data)
    if error:
        return False, error, None
    primary_billing = billing_employers[0] if billing_employers else None
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。", None
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, mail_no FROM mail_records
                WHERE created_by_user_id = %s
                  AND recipient_address_snapshot = %s
                  AND item_description = %s
                  AND created_at >= NOW() - INTERVAL '10 minutes'
                  AND status <> 'CANCELLED'
                ORDER BY id DESC LIMIT 1
                """,
                (user_id, address, description),
            )
            duplicate = cursor.fetchone()
            if duplicate and not data.get("confirm_duplicate"):
                return False, f"疑似重複：最近已建立 {duplicate['mail_no']}。", "DUPLICATE"

            cursor.execute(
                """
                INSERT INTO mail_records (
                    created_by_user_id, operator_name_snapshot,
                    sender_type, sender_employer_workplace_id, sender_name_snapshot,
                    sender_postal_code_snapshot, sender_address_snapshot,
                    sender_contact_snapshot, sender_phone_snapshot,
                    recipient_type,
                    employer_workplace_id, government_agency_id,
                    recipient_name_snapshot, recipient_department_snapshot,
                    recipient_postal_code_snapshot, recipient_address_snapshot,
                    recipient_contact_snapshot, billing_employer_workplace_id,
                    billing_employer_id_snapshot, billing_employer_name_snapshot,
                    item_description, mail_service, planned_shipping_fee, notes, status
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, 'READY'
                ) RETURNING id
                """,
                (
                    user_id,
                    operator_name,
                    sender["sender_type"],
                    sender["sender_employer_workplace_id"],
                    sender["sender_name"],
                    sender["sender_postal_code"],
                    sender["sender_address"],
                    sender["sender_contact"],
                    sender["sender_phone"],
                    data["recipient_type"],
                    data.get("employer_workplace_id"),
                    data.get("government_agency_id"),
                    name,
                    _clean(data.get("recipient_department")),
                    _clean(data.get("recipient_postal_code")),
                    address,
                    _clean(data.get("recipient_contact")),
                    primary_billing["employer_workplace_id"] if primary_billing else None,
                    primary_billing["employer_id"] if primary_billing else None,
                    primary_billing["employer_name"] if primary_billing else None,
                    description,
                    _clean(data.get("mail_service")) or "掛號",
                    planned_shipping_fee,
                    _clean(data.get("notes")),
                ),
            )
            record_id = cursor.fetchone()["id"]
            _replace_billing_employers(cursor, record_id, billing_employers, delete_existing=False)
            mail_no = f"MAIL-{datetime.now():%Y%m%d}-{record_id:06d}"
            cursor.execute("UPDATE mail_records SET mail_no = %s WHERE id = %s", (mail_no, record_id))
            _audit(cursor, "MAIL_RECORD", record_id, "CREATE", user_id, after={"mail_no": mail_no})
        conn.commit()
        return True, f"寄件已建立：{mail_no}", record_id
    except Exception as exc:
        conn.rollback()
        return False, f"建立寄件失敗：{exc}", None
    finally:
        conn.close()


def create_mail_records(user_id, operator_name, rows, confirm_duplicates=False, source="RECIPIENT_LIST"):
    """Atomically create multiple independent mail records."""
    prepared = []
    for index, data in enumerate(rows or [], 1):
        name = _clean(data.get("recipient_name"))
        address = _clean(data.get("recipient_address"))
        description = _clean(data.get("item_description"))
        if not name or not address or not description:
            return False, f"第 {index} 筆的收件人、地址與寄送內容皆為必填。", []
        recipient_type = data.get("recipient_type") or "CUSTOM"
        if recipient_type not in {"EMPLOYER", "GOVERNMENT", "CUSTOM"}:
            return False, f"第 {index} 筆的收件人類型不正確。", []
        sender, error = _normalize_sender(data)
        if error:
            return False, error, []
        try:
            planned_shipping_fee = _parse_shipping_fee(data.get("planned_shipping_fee"))
        except ValueError as exc:
            return False, f"第 {index} 筆：{exc}", []
        billing_employers, error = _normalize_billing_employers(data)
        if error:
            return False, f"第 {index} 筆：{error}", []
        prepared.append(
            {
                "data": data,
                "name": name,
                "address": address,
                "description": description,
                "recipient_type": recipient_type,
                "sender": sender,
                "planned_shipping_fee": planned_shipping_fee,
                "billing_employers": billing_employers,
            }
        )
    if not prepared:
        return False, "請至少選擇一位收件人。", []

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。", []
    try:
        with conn.cursor() as cursor:
            if not confirm_duplicates:
                for item in prepared:
                    cursor.execute(
                        """
                        SELECT mail_no FROM mail_records
                        WHERE created_by_user_id = %s
                          AND recipient_address_snapshot = %s
                          AND item_description = %s
                          AND created_at >= NOW() - INTERVAL '10 minutes'
                          AND status <> 'CANCELLED'
                        ORDER BY id DESC LIMIT 1
                        """,
                        (user_id, item["address"], item["description"]),
                    )
                    duplicate = cursor.fetchone()
                    if duplicate:
                        raise ValueError(
                            f"{item['name']} 可能重複建立（{duplicate['mail_no']}）；"
                            "請取消該筆，或勾選忽略重複提醒後再建立。"
                        )

            record_ids = []
            for item in prepared:
                data = item["data"]
                sender = item["sender"]
                billing_employers = item["billing_employers"]
                primary_billing = billing_employers[0] if billing_employers else None
                cursor.execute(
                    """
                    INSERT INTO mail_records (
                        created_by_user_id, operator_name_snapshot,
                        sender_type, sender_employer_workplace_id, sender_name_snapshot,
                        sender_postal_code_snapshot, sender_address_snapshot,
                        sender_contact_snapshot, sender_phone_snapshot,
                        recipient_type, employer_workplace_id, government_agency_id,
                        recipient_name_snapshot, recipient_department_snapshot,
                        recipient_postal_code_snapshot, recipient_address_snapshot,
                        recipient_contact_snapshot, billing_employer_workplace_id,
                        billing_employer_id_snapshot, billing_employer_name_snapshot,
                        item_description, mail_service, planned_shipping_fee, notes, status
                    ) VALUES (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'READY'
                    ) RETURNING id
                    """,
                    (
                        user_id, operator_name, sender["sender_type"],
                        sender["sender_employer_workplace_id"], sender["sender_name"],
                        sender["sender_postal_code"], sender["sender_address"],
                        sender["sender_contact"], sender["sender_phone"],
                        item["recipient_type"], data.get("employer_workplace_id"),
                        data.get("government_agency_id"), item["name"],
                        _clean(data.get("recipient_department")),
                        _clean(data.get("recipient_postal_code")), item["address"],
                        _clean(data.get("recipient_contact")),
                        primary_billing["employer_workplace_id"] if primary_billing else None,
                        primary_billing["employer_id"] if primary_billing else None,
                        primary_billing["employer_name"] if primary_billing else None,
                        item["description"], _clean(data.get("mail_service")) or "掛號",
                        item["planned_shipping_fee"], _clean(data.get("notes")),
                    ),
                )
                record_id = cursor.fetchone()["id"]
                _replace_billing_employers(cursor, record_id, billing_employers, delete_existing=False)
                mail_no = f"MAIL-{datetime.now():%Y%m%d}-{record_id:06d}"
                cursor.execute("UPDATE mail_records SET mail_no = %s WHERE id = %s", (mail_no, record_id))
                _audit(cursor, "MAIL_RECORD", record_id, "CREATE", user_id, after={"mail_no": mail_no, "source": source})
                record_ids.append(record_id)
        conn.commit()
        return True, f"已建立 {len(record_ids)} 筆寄件。", record_ids
    except Exception as exc:
        conn.rollback()
        return False, f"批次建立寄件失敗：{exc}", []
    finally:
        conn.close()


EDITABLE_STATUSES = ("DRAFT", "READY", "BATCHED")
_EDIT_AUDIT_FIELDS = (
    "sender_type", "sender_name_snapshot", "sender_postal_code_snapshot", "sender_address_snapshot",
    "sender_contact_snapshot", "sender_phone_snapshot", "recipient_name_snapshot",
    "recipient_department_snapshot", "recipient_postal_code_snapshot", "recipient_address_snapshot",
    "recipient_contact_snapshot", "billing_employer_name_snapshot", "item_description",
    "mail_service", "planned_shipping_fee", "notes",
)


def update_mail_record(record_id, user_id, data, allow_any_owner=False):
    """修改尚未寄出的寄件；總務完成寄件（DISPATCHED）後即不可再改。"""
    name = _clean(data.get("recipient_name"))
    address = _clean(data.get("recipient_address"))
    description = _clean(data.get("item_description"))
    if not name or not address or not description:
        return False, "收件人、地址與寄送內容皆為必填。"
    sender, error = _normalize_sender(data)
    if error:
        return False, error
    try:
        planned_shipping_fee = _parse_shipping_fee(data.get("planned_shipping_fee"))
    except ValueError as exc:
        return False, str(exc)
    billing_employers, error = _normalize_billing_employers(data)
    if error:
        return False, error
    primary_billing = billing_employers[0] if billing_employers else None
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT r.*,
                       COALESCE(
                           (
                               SELECT JSONB_AGG(
                                   JSONB_BUILD_OBJECT(
                                       'key', be.employer_key,
                                       'employer_workplace_id', be.employer_workplace_id,
                                       'employer_id', be.employer_id_snapshot,
                                       'employer_name', be.employer_name_snapshot
                                   ) ORDER BY be.sequence_no
                               )
                               FROM mail_record_billing_employers be
                               WHERE be.mail_record_id = r.id
                           ),
                           '[]'::JSONB
                       ) AS billing_employers
                FROM mail_records r
                WHERE r.id = %s
                FOR UPDATE
                """,
                (record_id,),
            )
            before = cursor.fetchone()
            if not before:
                return False, "找不到寄件資料。"
            if not allow_any_owner and before["created_by_user_id"] != user_id:
                return False, "只能修改自己建立的寄件。"
            if before["status"] not in EDITABLE_STATUSES:
                return False, "此寄件已寄出或已取消，不可再修改。"
            cursor.execute(
                """
                UPDATE mail_records SET
                    sender_type = %s, sender_employer_workplace_id = %s, sender_name_snapshot = %s,
                    sender_postal_code_snapshot = %s, sender_address_snapshot = %s,
                    sender_contact_snapshot = %s, sender_phone_snapshot = %s,
                    recipient_name_snapshot = %s, recipient_department_snapshot = %s,
                    recipient_postal_code_snapshot = %s, recipient_address_snapshot = %s,
                    recipient_contact_snapshot = %s, billing_employer_workplace_id = %s,
                    billing_employer_id_snapshot = %s, billing_employer_name_snapshot = %s,
                    item_description = %s, mail_service = %s, planned_shipping_fee = %s,
                    notes = %s, updated_at = NOW()
                WHERE id = %s AND status IN %s
                RETURNING *
                """,
                (
                    sender["sender_type"],
                    sender["sender_employer_workplace_id"],
                    sender["sender_name"],
                    sender["sender_postal_code"],
                    sender["sender_address"],
                    sender["sender_contact"],
                    sender["sender_phone"],
                    name,
                    _clean(data.get("recipient_department")),
                    _clean(data.get("recipient_postal_code")),
                    address,
                    _clean(data.get("recipient_contact")),
                    primary_billing["employer_workplace_id"] if primary_billing else None,
                    primary_billing["employer_id"] if primary_billing else None,
                    primary_billing["employer_name"] if primary_billing else None,
                    description,
                    _clean(data.get("mail_service")) or "掛號",
                    planned_shipping_fee,
                    _clean(data.get("notes")),
                    record_id,
                    EDITABLE_STATUSES,
                ),
            )
            after = cursor.fetchone()
            _replace_billing_employers(cursor, record_id, billing_employers)
            changed = [field for field in _EDIT_AUDIT_FIELDS if before.get(field) != after.get(field)]
            before_billing = [
                {
                    "employer_id": _clean(item.get("employer_id")),
                    "employer_name": _clean(item.get("employer_name")),
                }
                for item in (before.get("billing_employers") or [])
            ]
            after_billing = [
                {
                    "employer_id": item["employer_id"],
                    "employer_name": item["employer_name"],
                }
                for item in billing_employers
            ]
            if before_billing != after_billing:
                changed.append("billing_employers")
            if changed:
                before_audit = {
                    field: (before_billing if field == "billing_employers" else before.get(field))
                    for field in changed
                }
                after_audit = {
                    field: (after_billing if field == "billing_employers" else after.get(field))
                    for field in changed
                }
                _audit(
                    cursor,
                    "MAIL_RECORD",
                    record_id,
                    "UPDATE",
                    user_id,
                    before=before_audit,
                    after=after_audit,
                )
        conn.commit()
        return True, f"已更新 {before['mail_no']}。"
    except Exception as exc:
        conn.rollback()
        return False, f"修改寄件失敗：{exc}"
    finally:
        conn.close()


def get_mail_record(record_id):
    rows = _fetch_all(
        """
        SELECT r.*,
               COALESCE(
                   (
                       SELECT JSONB_AGG(
                           JSONB_BUILD_OBJECT(
                               'key', be.employer_key,
                               'employer_workplace_id', be.employer_workplace_id,
                               'employer_id', be.employer_id_snapshot,
                               'employer_name', be.employer_name_snapshot
                           ) ORDER BY be.sequence_no
                       )
                       FROM mail_record_billing_employers be
                       WHERE be.mail_record_id = r.id
                   ),
                   CASE
                       WHEN r.billing_employer_name_snapshot IS NOT NULL THEN
                           JSONB_BUILD_ARRAY(
                               JSONB_BUILD_OBJECT(
                                   'key', CASE
                                       WHEN r.billing_employer_id_snapshot IS NOT NULL
                                           THEN 'CODE:' || LOWER(r.billing_employer_id_snapshot)
                                       ELSE 'NAME:' || LOWER(r.billing_employer_name_snapshot)
                                   END,
                                   'employer_workplace_id', r.billing_employer_workplace_id,
                                   'employer_id', r.billing_employer_id_snapshot,
                                   'employer_name', r.billing_employer_name_snapshot
                               )
                           )
                       ELSE '[]'::JSONB
                   END
               ) AS billing_employers
        FROM mail_records r
        WHERE r.id = %s
        """,
        (record_id,),
    )
    return rows[0] if rows else None


def get_user_mail_records(user_id, keyword=None, status=None, limit=100):
    clauses = ["r.created_by_user_id = %s"]
    params = [user_id]
    if keyword:
        clauses.append(
            "(r.recipient_name_snapshot ILIKE %s OR r.sender_name_snapshot ILIKE %s OR "
            "r.item_description ILIKE %s OR b.batch_no ILIKE %s OR bi.tracking_no ILIKE %s OR "
            "COALESCE((SELECT STRING_AGG(CONCAT_WS(' ', be.employer_id_snapshot, "
            "be.employer_name_snapshot), ' ') FROM mail_record_billing_employers be "
            "WHERE be.mail_record_id = r.id), CONCAT_WS(' ', r.billing_employer_id_snapshot, "
            "r.billing_employer_name_snapshot)) ILIKE %s)"
        )
        term = f"%{keyword}%"
        params.extend([term] * 6)
    if status:
        clauses.append("r.status = %s")
        params.append(status)
    params.append(limit)
    return _fetch_all(
        f"""
        SELECT r.id, r.mail_no, r.created_at, r.sender_type, r.sender_name_snapshot,
               r.recipient_name_snapshot, r.recipient_department_snapshot,
               r.recipient_contact_snapshot, r.recipient_postal_code_snapshot,
               r.recipient_address_snapshot,
               r.item_description, r.mail_service, r.notes,
               COALESCE(
                   (SELECT STRING_AGG(be.employer_name_snapshot, '、' ORDER BY be.sequence_no)
                    FROM mail_record_billing_employers be WHERE be.mail_record_id = r.id),
                   r.billing_employer_name_snapshot,
                   ''
               ) AS billing_employer_names,
               r.status, r.planned_shipping_fee, b.batch_no, b.posting_date,
               bi.tracking_no, bi.shipping_fee, b.id AS batch_id,
               b.receipt_file_path, b.receipt_original_name, b.receipt_content_type,
               dispatcher.display_name AS processor_name
        FROM mail_records r
        LEFT JOIN mail_batch_items bi ON bi.mail_record_id = r.id AND bi.removed_at IS NULL
        LEFT JOIN mail_batches b ON b.id = bi.batch_id AND b.status <> 'VOID'
        LEFT JOIN mail_users dispatcher ON dispatcher.id = b.dispatched_by_user_id
        WHERE {' AND '.join(clauses)}
        ORDER BY r.created_at DESC LIMIT %s
        """,
        tuple(params),
    )


def cancel_mail_record(record_id, user_id, reason, allow_any_owner=False):
    """取消尚未交寄的寄件；郵局作業可代為取消任何人建立的寄件。"""
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            owner_clause = "" if allow_any_owner else " AND created_by_user_id = %s"
            params = [user_id, _clean(reason) or "使用者取消", record_id]
            if not allow_any_owner:
                params.append(user_id)
            cursor.execute(
                f"""
                UPDATE mail_records SET status = 'CANCELLED', cancelled_at = NOW(),
                    cancelled_by_user_id = %s, cancel_reason = %s, updated_at = NOW()
                WHERE id = %s{owner_clause} AND status IN ('DRAFT', 'READY')
                RETURNING mail_no
                """,
                tuple(params),
            )
            row = cursor.fetchone()
            if not row:
                return False, "找不到可取消的寄件；此寄件可能已進入批次、已寄出或已取消。"
            _audit(cursor, "MAIL_RECORD", record_id, "CANCEL", user_id, after={"reason": reason})
        conn.commit()
        return True, f"已取消 {row['mail_no']}。"
    except Exception as exc:
        conn.rollback()
        return False, f"取消失敗：{exc}"
    finally:
        conn.close()


def restore_cancelled_mail_record(record_id, user_id):
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE mail_records SET status = 'READY', cancelled_at = NULL,
                    cancelled_by_user_id = NULL, cancel_reason = NULL, updated_at = NOW()
                WHERE id = %s AND created_by_user_id = %s AND status = 'CANCELLED'
                RETURNING mail_no
                """,
                (record_id, user_id),
            )
            row = cursor.fetchone()
            if not row:
                return False, "只能恢復自己已取消的寄件。"
            _audit(
                cursor,
                "MAIL_RECORD",
                record_id,
                "RESTORE_READY",
                user_id,
                before={"status": "CANCELLED"},
                after={"status": "READY"},
            )
        conn.commit()
        return True, f"{row['mail_no']} 已恢復為待寄出。"
    except Exception as exc:
        conn.rollback()
        return False, f"恢復待寄出失敗：{exc}"
    finally:
        conn.close()


# ---------- 批次 ----------

def get_pending_records():
    return _fetch_all(
        """
        SELECT r.id, r.mail_no, r.created_at, u.display_name AS operator_name,
               r.sender_type, r.sender_name_snapshot,
               r.recipient_name_snapshot, r.recipient_department_snapshot,
               r.recipient_contact_snapshot, r.recipient_postal_code_snapshot,
               r.recipient_address_snapshot, r.item_description, r.mail_service,
               r.planned_sequence, r.planned_shipping_fee, r.notes
        FROM mail_records r
        JOIN mail_users u ON u.id = r.created_by_user_id
        WHERE r.status = 'READY'
        ORDER BY r.created_at, r.id
        """
    )


def save_pending_postage(item_fees, user_id, item_sequences=None):
    normalized_fees = {}
    normalized_sequences = {}
    update_sequences = item_sequences is not None
    try:
        for record_id, raw_fee in (item_fees or {}).items():
            normalized_fees[int(record_id)] = _parse_shipping_fee(raw_fee)
        for record_id, raw_sequence in (item_sequences or {}).items():
            record_id = int(record_id)
            try:
                sequence = _parse_shipping_fee(raw_sequence)
            except (TypeError, ValueError, OverflowError) as exc:
                raise ValueError("順序必須是正整數。") from exc
            if sequence is not None and sequence < 1:
                raise ValueError("順序必須從 1 開始。")
            normalized_sequences[record_id] = sequence
    except (ValueError, TypeError) as exc:
        return False, str(exc)
    record_ids = list(dict.fromkeys([*normalized_fees, *normalized_sequences]))
    if not record_ids:
        return False, "目前沒有可儲存的待寄資料。"

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT id, status FROM mail_records WHERE id = ANY(%s) FOR UPDATE",
                (record_ids,),
            )
            rows = {row["id"]: row["status"] for row in cursor.fetchall()}
            if set(rows) != set(record_ids) or any(status != "READY" for status in rows.values()):
                raise ValueError("部分寄件已被其他人納入批次，請重新整理後再試。")
            for record_id in record_ids:
                if update_sequences:
                    cursor.execute(
                        """
                        UPDATE mail_records
                        SET planned_sequence = %s, planned_shipping_fee = %s, updated_at = NOW()
                        WHERE id = %s
                        """,
                        (
                            normalized_sequences.get(record_id),
                            normalized_fees.get(record_id),
                            record_id,
                        ),
                    )
                else:
                    cursor.execute(
                        """
                        UPDATE mail_records
                        SET planned_shipping_fee = %s, updated_at = NOW()
                        WHERE id = %s
                        """,
                        (normalized_fees.get(record_id), record_id),
                    )
            _audit(
                cursor,
                "MAIL_RECORD",
                0,
                "SAVE_PLANNED_POSTAGE",
                user_id,
                after={"record_ids": record_ids, "count": len(record_ids)},
            )
        conn.commit()
        return True, f"已儲存 {len(record_ids)} 筆待寄順序與郵資。"
    except Exception as exc:
        conn.rollback()
        return False, f"儲存順序與郵資失敗：{exc}"
    finally:
        conn.close()


def _normalize_posting_date(value):
    """交寄日期：預設今天；可補登過去日期，但不可是未來。"""
    if value is None:
        return date.today()
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        value = date.fromisoformat(str(value))
    if value > date.today():
        raise ValueError("交寄日期不可晚於今天。")
    return value


def generate_sequential_tracking_numbers(start, suffix, count):
    """依流水碼起始值與固定後碼產生一批掛號號碼，並保留前導零。"""
    start = str(start or "").strip()
    suffix = str(suffix or "").strip()
    if not start or not suffix:
        raise ValueError("請填寫流水碼起始與固定後碼。")
    if not start.isdigit() or not suffix.isdigit():
        raise ValueError("流水碼起始與固定後碼只能輸入數字。")
    count = int(count)
    if count < 1:
        raise ValueError("此批次沒有可編號的寄件。")

    width = len(start)
    first = int(start)
    last = first + count - 1
    if len(str(last)) > width:
        raise ValueError(f"流水碼會超過 {width} 碼，請調整起始值。")
    return [f"{number:0{width}d}{suffix}" for number in range(first, last + 1)]


def create_batch(record_ids, submission_mode, user_id, item_fees=None, posting_date=None):
    # 保留畫面拖曳後的順序；dict.fromkeys 同時避免同一寄件被重複加入。
    ids = list(dict.fromkeys(int(value) for value in record_ids))
    if not ids:
        return False, "請至少選擇一筆待寄件。", None
    try:
        posting_date = _normalize_posting_date(posting_date)
    except ValueError as exc:
        return False, str(exc), None
    if submission_mode not in {"BULK_FORM", "REGULAR"}:
        return False, "交寄方式不正確。", None
    template_name = "郵局大宗掛號單.xlsx" if submission_mode == "BULK_FORM" else None
    if template_name and not (TEMPLATE_DIR / template_name).exists():
        return False, f"找不到範本：doc_templates/{template_name}", None
    normalized_fees = {}
    for record_id, raw_fee in (item_fees or {}).items():
        try:
            normalized_fees[int(record_id)] = _parse_shipping_fee(raw_fee)
        except ValueError as exc:
            return False, str(exc), None
    entered_fees = [fee for record_id, fee in normalized_fees.items() if record_id in ids and fee is not None]
    initial_total_fee = sum(entered_fees) if entered_fees else None

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。", None
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT id FROM mail_records WHERE id = ANY(%s) AND status = 'READY' FOR UPDATE",
                (ids,),
            )
            locked = [row["id"] for row in cursor.fetchall()]
            if set(locked) != set(ids):
                raise ValueError("部分寄件已被其他人納入批次，請重新整理。")
            template_version = None
            if template_name:
                template_version = hashlib.sha256((TEMPLATE_DIR / template_name).read_bytes()).hexdigest()[:16]
            cursor.execute(
                """
                INSERT INTO mail_batches
                    (submission_mode, template_filename, template_version,
                     total_shipping_fee, created_by_user_id, posting_date)
                VALUES (%s, %s, %s, %s, %s, %s) RETURNING id
                """,
                (submission_mode, template_name, template_version, initial_total_fee, user_id, posting_date),
            )
            batch_id = cursor.fetchone()["id"]
            batch_no = f"POST-{posting_date:%Y%m%d}-{batch_id:05d}"
            cursor.execute("UPDATE mail_batches SET batch_no = %s WHERE id = %s", (batch_no, batch_id))
            for sequence, record_id in enumerate(ids, 1):
                cursor.execute(
                    """
                    INSERT INTO mail_batch_items
                        (batch_id, mail_record_id, sequence_no, shipping_fee)
                    VALUES (%s, %s, %s, %s)
                    """,
                    (batch_id, record_id, sequence, normalized_fees.get(record_id)),
                )
            cursor.execute(
                "UPDATE mail_records SET status = 'BATCHED', updated_at = NOW() WHERE id = ANY(%s)",
                (ids,),
            )
            _audit(
                cursor,
                "MAIL_BATCH",
                batch_id,
                "CREATE",
                user_id,
                after={"batch_no": batch_no, "record_ids": ids, "posting_date": posting_date.isoformat()},
            )
        conn.commit()
        return True, f"已建立批次 {batch_no}（{len(ids)} 件）。", batch_id
    except Exception as exc:
        conn.rollback()
        return False, f"建立批次失敗：{exc}", None
    finally:
        conn.close()


def update_batch_submission_mode(batch_id, submission_mode, user_id):
    if submission_mode not in {"BULK_FORM", "REGULAR"}:
        return False, "交寄方式不正確。"
    template_name = "郵局大宗掛號單.xlsx" if submission_mode == "BULK_FORM" else None
    if template_name and not (TEMPLATE_DIR / template_name).exists():
        return False, f"找不到範本：doc_templates/{template_name}"
    template_version = (
        hashlib.sha256((TEMPLATE_DIR / template_name).read_bytes()).hexdigest()[:16]
        if template_name
        else None
    )

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT id, batch_no, status, submission_mode FROM mail_batches WHERE id = %s FOR UPDATE",
                (batch_id,),
            )
            batch = cursor.fetchone()
            if not batch or batch["status"] != "CREATED":
                return False, "只有等待結案的批次可以更正交寄方式。"
            if batch["submission_mode"] == submission_mode:
                return True, "交寄方式未變更。"
            cursor.execute(
                """
                UPDATE mail_batches
                SET submission_mode = %s, template_filename = %s, template_version = %s
                WHERE id = %s
                """,
                (submission_mode, template_name, template_version, batch_id),
            )
            _audit(
                cursor,
                "MAIL_BATCH",
                batch_id,
                "CHANGE_SUBMISSION_MODE",
                user_id,
                before={"submission_mode": batch["submission_mode"]},
                after={"submission_mode": submission_mode},
            )
        conn.commit()
        return True, f"已將 {batch['batch_no']} 改為{('大宗單' if submission_mode == 'BULK_FORM' else '一般交寄')}。"
    except Exception as exc:
        conn.rollback()
        return False, f"更正交寄方式失敗：{exc}"
    finally:
        conn.close()


def get_batches(status=None, limit=200):
    where = "WHERE b.status = %s" if status else ""
    params = (status, limit) if status else (limit,)
    return _fetch_all(
        f"""
        SELECT b.*, u.display_name AS created_by_name,
               COUNT(bi.id) FILTER (WHERE bi.removed_at IS NULL) AS item_count
        FROM mail_batches b
        JOIN mail_users u ON u.id = b.created_by_user_id
        LEFT JOIN mail_batch_items bi ON bi.batch_id = b.id
        {where}
        GROUP BY b.id, u.display_name
        ORDER BY b.created_at DESC LIMIT %s
        """,
        params,
    )


def get_batch(batch_id):
    batches = _fetch_all("SELECT * FROM mail_batches WHERE id = %s", (batch_id,))
    if not batches:
        return None, []
    items = _fetch_all(
        """
        SELECT bi.id AS batch_item_id, bi.sequence_no, bi.tracking_no, bi.shipping_fee,
               r.id AS mail_record_id, r.mail_no, r.recipient_name_snapshot,
               r.recipient_postal_code_snapshot, r.recipient_address_snapshot,
               r.item_description, r.mail_service, r.notes,
               u.display_name AS operator_name,
               r.sender_type, r.sender_name_snapshot
        FROM mail_batch_items bi
        JOIN mail_records r ON r.id = bi.mail_record_id
        JOIN mail_users u ON u.id = r.created_by_user_id
        WHERE bi.batch_id = %s AND bi.removed_at IS NULL
        ORDER BY bi.sequence_no
        """,
        (batch_id,),
    )
    return batches[0], items


def save_batch_tracking_numbers(batch_id, user_id, tracking_numbers):
    numbers = [_clean(value) for value in tracking_numbers]
    if not numbers or any(not value or not value.isdigit() for value in numbers):
        return False, "掛號號碼只能包含數字，且不可空白。"
    if len(set(numbers)) != len(numbers):
        return False, "掛號號碼不可重複。"

    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT status FROM mail_batches WHERE id = %s FOR UPDATE", (batch_id,))
            batch = cursor.fetchone()
            if not batch or batch["status"] != "CREATED":
                raise ValueError("此批次不存在或已經結案。")
            cursor.execute(
                """
                SELECT id FROM mail_batch_items
                WHERE batch_id = %s AND removed_at IS NULL
                ORDER BY sequence_no
                FOR UPDATE
                """,
                (batch_id,),
            )
            item_ids = [row["id"] for row in cursor.fetchall()]
            if len(item_ids) != len(numbers):
                raise ValueError("批次件數已變更，請重新整理後再試。")
            # 先清空同批舊號碼，避免調整起始值時被同批的唯一鍵暫時衝突。
            cursor.execute(
                "UPDATE mail_batch_items SET tracking_no = NULL WHERE batch_id = %s AND removed_at IS NULL",
                (batch_id,),
            )
            for item_id, tracking_number in zip(item_ids, numbers):
                cursor.execute(
                    "UPDATE mail_batch_items SET tracking_no = %s WHERE id = %s",
                    (tracking_number, item_id),
                )
            _audit(
                cursor,
                "MAIL_BATCH",
                batch_id,
                "SAVE_TRACKING_NUMBERS",
                user_id,
                after={"count": len(numbers), "first": numbers[0], "last": numbers[-1]},
            )
        conn.commit()
        return True, f"已更新 {len(numbers)} 筆掛號號碼（{numbers[0]} ～ {numbers[-1]}）。"
    except Exception as exc:
        conn.rollback()
        return False, f"儲存掛號號碼失敗：{exc}"
    finally:
        conn.close()


def complete_batch(batch_id, user_id, total_fee, item_updates, receipt_meta, posting_date=None):
    try:
        posting_date = _normalize_posting_date(posting_date) if posting_date is not None else None
    except ValueError as exc:
        return False, str(exc)
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT * FROM mail_batches WHERE id = %s FOR UPDATE", (batch_id,))
            batch = cursor.fetchone()
            if not batch or batch["status"] != "CREATED":
                return False, "此批次不存在或已經結案。"
            for update in item_updates or []:
                tracking = _clean(update.get("tracking_no"))
                fee = _parse_shipping_fee(update.get("shipping_fee"))
                cursor.execute(
                    """
                    UPDATE mail_batch_items SET tracking_no = %s, shipping_fee = %s
                    WHERE id = %s AND batch_id = %s AND removed_at IS NULL
                    """,
                    (tracking, fee, update["batch_item_id"], batch_id),
                )
            cursor.execute(
                """
                UPDATE mail_batches SET status = 'DISPATCHED', total_shipping_fee = %s,
                    receipt_file_path = %s, receipt_original_name = %s,
                    receipt_content_type = %s, receipt_uploaded_at = NOW(),
                    receipt_uploaded_by_user_id = %s, dispatched_at = NOW(),
                    dispatched_by_user_id = %s,
                    posting_date = COALESCE(%s, posting_date, CURRENT_DATE)
                WHERE id = %s
                """,
                (
                    _parse_shipping_fee(total_fee),
                    receipt_meta["path"],
                    receipt_meta["original_name"],
                    receipt_meta["content_type"],
                    user_id,
                    user_id,
                    posting_date,
                    batch_id,
                ),
            )
            cursor.execute(
                """
                UPDATE mail_records SET status = 'DISPATCHED', updated_at = NOW()
                WHERE id IN (
                    SELECT mail_record_id FROM mail_batch_items
                    WHERE batch_id = %s AND removed_at IS NULL
                )
                """,
                (batch_id,),
            )
            _audit(
                cursor, "MAIL_BATCH", batch_id, "DISPATCH", user_id,
                after={"total_fee": str(total_fee), "posting_date": posting_date.isoformat() if posting_date else None},
            )
        conn.commit()
        return True, "批次已完成寄出。"
    except Exception as exc:
        conn.rollback()
        return False, f"批次結案失敗：{exc}"
    finally:
        conn.close()


def update_dispatched_batch(
    batch_id, user_id, total_fee, item_updates, posting_date, receipt_meta=None
):
    """更新已結案批次的交寄資料；未上傳新收據時保留原收據。"""
    try:
        posting_date = _normalize_posting_date(posting_date)
        total_fee = _parse_shipping_fee(total_fee)
    except ValueError as exc:
        return False, str(exc)
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT * FROM mail_batches WHERE id = %s FOR UPDATE", (batch_id,))
            before = cursor.fetchone()
            if not before or before["status"] != "DISPATCHED":
                return False, "只有已結案的批次可以修改。"

            normalized_items = []
            for update in item_updates or []:
                normalized_items.append(
                    {
                        "batch_item_id": int(update["batch_item_id"]),
                        "tracking_no": _clean(update.get("tracking_no")),
                        "shipping_fee": _parse_shipping_fee(update.get("shipping_fee")),
                    }
                )
            for update in normalized_items:
                cursor.execute(
                    """
                    UPDATE mail_batch_items SET tracking_no = %s, shipping_fee = %s
                    WHERE id = %s AND batch_id = %s AND removed_at IS NULL
                    """,
                    (
                        update["tracking_no"],
                        update["shipping_fee"],
                        update["batch_item_id"],
                        batch_id,
                    ),
                )

            receipt_fields = ""
            params = [total_fee, posting_date]
            if receipt_meta:
                receipt_fields = """,
                    receipt_file_path = %s, receipt_original_name = %s,
                    receipt_content_type = %s, receipt_uploaded_at = NOW(),
                    receipt_uploaded_by_user_id = %s"""
                params.extend(
                    [
                        receipt_meta["path"],
                        receipt_meta["original_name"],
                        receipt_meta["content_type"],
                        user_id,
                    ]
                )
            params.append(batch_id)
            cursor.execute(
                f"""
                UPDATE mail_batches SET total_shipping_fee = %s, posting_date = %s
                    {receipt_fields}
                WHERE id = %s AND status = 'DISPATCHED'
                """,
                tuple(params),
            )
            audit_before = {
                "total_fee": (
                    str(before.get("total_shipping_fee"))
                    if before.get("total_shipping_fee") is not None
                    else None
                ),
                "posting_date": (
                    before["posting_date"].isoformat()
                    if hasattr(before.get("posting_date"), "isoformat")
                    else before.get("posting_date")
                ),
                "receipt_file_path": before.get("receipt_file_path"),
            }
            audit_after = {
                "total_fee": total_fee,
                "posting_date": posting_date.isoformat(),
                "receipt_file_path": (
                    receipt_meta["path"] if receipt_meta else before.get("receipt_file_path")
                ),
                "items": normalized_items,
            }
            _audit(
                cursor,
                "MAIL_BATCH",
                batch_id,
                "UPDATE_DISPATCHED",
                user_id,
                before=audit_before,
                after=audit_after,
            )
        conn.commit()
        return True, "已更新過去的結案紀錄。"
    except Exception as exc:
        conn.rollback()
        return False, f"修改結案紀錄失敗：{exc}"
    finally:
        conn.close()


def void_batch(batch_id, user_id, reason):
    reason = _clean(reason)
    if not reason:
        return False, "請填寫作廢原因。"
    conn = database.get_db_connection()
    if not conn:
        return False, "資料庫連線失敗。"
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT status FROM mail_batches WHERE id = %s FOR UPDATE", (batch_id,))
            row = cursor.fetchone()
            if not row or row["status"] != "CREATED":
                return False, "只有尚未結案的批次可以作廢。"
            cursor.execute(
                """
                UPDATE mail_records r
                SET status = 'READY',
                    planned_shipping_fee = COALESCE(bi.shipping_fee, r.planned_shipping_fee),
                    updated_at = NOW()
                FROM mail_batch_items bi
                WHERE bi.mail_record_id = r.id
                  AND bi.batch_id = %s
                  AND bi.removed_at IS NULL
                """,
                (batch_id,),
            )
            cursor.execute(
                """
                UPDATE mail_batch_items SET removed_at = NOW(), removed_by_user_id = %s,
                    remove_reason = %s WHERE batch_id = %s AND removed_at IS NULL
                """,
                (user_id, reason, batch_id),
            )
            cursor.execute(
                """
                UPDATE mail_batches SET status = 'VOID', voided_at = NOW(),
                    voided_by_user_id = %s, void_reason = %s WHERE id = %s
                """,
                (user_id, reason, batch_id),
            )
            _audit(cursor, "MAIL_BATCH", batch_id, "VOID", user_id, after={"reason": reason})
        conn.commit()
        return True, "批次已作廢，信件已退回待寄。"
    except Exception as exc:
        conn.rollback()
        return False, f"作廢失敗：{exc}"
    finally:
        conn.close()


def search_all_records(keyword=None, status=None, limit=500):
    clauses = ["1 = 1"]
    params = []
    if keyword:
        clauses.append(
            "(r.mail_no ILIKE %s OR r.recipient_name_snapshot ILIKE %s OR "
            "r.operator_name_snapshot ILIKE %s OR r.sender_name_snapshot ILIKE %s OR "
            "r.item_description ILIKE %s OR b.batch_no ILIKE %s OR bi.tracking_no ILIKE %s OR "
            "COALESCE((SELECT STRING_AGG(CONCAT_WS(' ', be.employer_id_snapshot, "
            "be.employer_name_snapshot), ' ') FROM mail_record_billing_employers be "
            "WHERE be.mail_record_id = r.id), CONCAT_WS(' ', r.billing_employer_id_snapshot, "
            "r.billing_employer_name_snapshot)) ILIKE %s)"
        )
        term = f"%{keyword}%"
        params.extend([term] * 8)
    if status:
        clauses.append("r.status = %s")
        params.append(status)
    params.append(limit)
    return _fetch_all(
        f"""
        SELECT r.id, r.mail_no, r.created_at, r.operator_name_snapshot,
               dispatcher.display_name AS processor_name,
               r.sender_type, r.sender_name_snapshot,
               r.recipient_name_snapshot, r.recipient_department_snapshot,
               r.recipient_contact_snapshot, r.recipient_address_snapshot,
               r.item_description, r.mail_service, r.notes,
               COALESCE(
                   (SELECT STRING_AGG(be.employer_name_snapshot, '、' ORDER BY be.sequence_no)
                    FROM mail_record_billing_employers be WHERE be.mail_record_id = r.id),
                   r.billing_employer_name_snapshot,
                   ''
               ) AS billing_employer_names,
               r.status,
               b.id AS batch_id, b.batch_no, b.submission_mode,
               b.total_shipping_fee, b.dispatched_at, b.posting_date, bi.tracking_no, bi.shipping_fee,
               b.receipt_file_path, b.receipt_original_name, b.receipt_content_type
        FROM mail_records r
        LEFT JOIN mail_batch_items bi ON bi.mail_record_id = r.id AND bi.removed_at IS NULL
        LEFT JOIN mail_batches b ON b.id = bi.batch_id AND b.status <> 'VOID'
        LEFT JOIN mail_users dispatcher ON dispatcher.id = b.dispatched_by_user_id
        WHERE {' AND '.join(clauses)}
        ORDER BY r.created_at DESC LIMIT %s
        """,
        tuple(params),
    )


# ---------- 文件與附件 ----------

def generate_bulk_workbook(batch_id):
    batch, items = get_batch(batch_id)
    if not batch or batch["submission_mode"] != "BULK_FORM":
        raise ValueError("此批次未使用郵局大宗單。")
    template_path = TEMPLATE_DIR / (batch.get("template_filename") or "郵局大宗掛號單.xlsx")
    if not template_path.exists():
        raise FileNotFoundError(f"找不到範本：{template_path.name}")
    workbook = load_workbook(template_path)
    template_sheet = workbook[workbook.sheetnames[0]]
    chunks = [items[index:index + 40] for index in range(0, len(items), 40)] or [[]]
    sheets = [template_sheet]
    for _ in range(1, len(chunks)):
        sheets.append(workbook.copy_worksheet(template_sheet))
    posted_on = batch.get("posting_date") or date.today()
    for index, (sheet, chunk) in enumerate(zip(sheets, chunks), 1):
        sheet.title = f"P{index}"
        sheet.print_area = "A1:J45"
        sheet.print_title_rows = "1:4"
        sheet["C2"] = posted_on.year - 1911
        sheet["E2"] = posted_on.month
        sheet["G2"] = posted_on.day
        for row in range(5, 45):
            for column in ("B", "D", "H", "J"):
                sheet[f"{column}{row}"] = None
        for offset, item in enumerate(chunk, 5):
            sheet[f"B{offset}"] = item.get("tracking_no") or ""
            sheet[f"D{offset}"] = item["recipient_name_snapshot"]
            note = "  ".join(
                str(value).strip()
                for value in (
                    resolved_recipient_postal_code(item),
                    item.get("item_description"),
                    item.get("notes"),
                )
                if value is not None and str(value).strip()
            )
            sheet[f"H{offset}"] = note
            sheet[f"J{offset}"] = item.get("shipping_fee")
        # Do not rely on Excel recalculating the template formula. Some browser
        # previews and alternative spreadsheet apps otherwise show a blank count.
        sheet["F45"] = "共"
        sheet["G45"] = len(chunk)
        sheet["I45"] = "件"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), f"郵局大宗掛號單_{batch['batch_no']}.xlsx"


def list_envelope_templates():
    if not TEMPLATE_DIR.exists():
        return []
    return sorted(path.name for path in TEMPLATE_DIR.glob("*.docx") if not path.name.startswith("~$"))


def _replace_docx_text(paragraph, mapping):
    original = paragraph.text
    replaced = original
    for key, value in mapping.items():
        replaced = replaced.replace(key, value or "")
    if replaced == original:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)


def _remove_docx_text_boxes(document):
    """Drop floating text boxes; the small envelope keeps a stale off-page copy of old data."""
    for container in list(document.element.body.iter(qn("w:pict"), qn("w:drawing"))):
        if next(container.iter(qn("w:txbxContent")), None) is not None:
            container.getparent().remove(container)


def _replace_docx_run_text(paragraph, old, new):
    """Replace text that may span runs without flattening the paragraph layout."""
    runs = paragraph.runs
    full_text = "".join(run.text for run in runs)
    start = full_text.find(old)
    if start < 0:
        return False

    end = start + len(old)
    position = 0
    first_index = first_offset = last_index = last_offset = None
    for index, run in enumerate(runs):
        run_end = position + len(run.text)
        if first_index is None and start < run_end:
            first_index = index
            first_offset = start - position
        if first_index is not None and end <= run_end:
            last_index = index
            last_offset = end - position
            break
        position = run_end

    if first_index is None or last_index is None:
        return False
    if first_index == last_index:
        text = runs[first_index].text
        runs[first_index].text = text[:first_offset] + (new or "") + text[last_offset:]
        return True

    runs[first_index].text = runs[first_index].text[:first_offset] + (new or "")
    for index in range(first_index + 1, last_index):
        runs[index].text = ""
    runs[last_index].text = runs[last_index].text[last_offset:]
    return True


def _is_company_recipient(record):
    return (
        _clean(record.get("recipient_name_snapshot")) == COMPANY_SENDER["name"]
        and _clean(record.get("recipient_address_snapshot")) == COMPANY_SENDER["address"]
    )


def _recipient_department_value(record):
    department = record.get("recipient_department_snapshot") or ""
    if _is_company_recipient(record):
        return " ".join(value for value in (department, COMPANY_SENDER["phone"]) if value)
    return department


def _recipient_department_contact(record):
    department = _recipient_department_value(record)
    contact = record.get("recipient_contact_snapshot") or ""
    return "：".join(value for value in (department, contact) if value)


def resolved_recipient_postal_code(record):
    """Use a saved code first, otherwise infer a non-persisted 3-digit default."""
    saved = _clean(
        record.get("recipient_postal_code_snapshot")
        or record.get("recipient_postal_code")
    )
    if saved:
        return saved
    address = record.get("recipient_address_snapshot") or record.get("recipient_address")
    return infer_postal_code(address) or ""


def _sender_display_value(record, snapshot_field, company_field):
    """Use fixed company details for older company records whose snapshots are blank."""
    value = record.get(snapshot_field)
    if value:
        return value
    if (record.get("sender_type") or "COMPANY") == "COMPANY":
        return COMPANY_SENDER[company_field]
    return ""


NAMED_ENVELOPE_LABELS = (
    "寄件人名稱",
    "寄件人郵遞區號",
    "寄件人郵寄區號",
    "寄件人地址",
    "寄件聯絡人",
    "寄件人電話",
    "郵遞區號",
    "收件地址",
    "寄送內容",
    "收件人／單位名稱",
    "收件部門：收件聯絡人",
    "收件部門",
    "收件聯絡人",
    "郵件種類",
)


def _docx_content_paragraphs(document):
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def _uses_named_envelope_labels(document):
    return any(
        label in paragraph.text
        for paragraph in _docx_content_paragraphs(document)
        for label in NAMED_ENVELOPE_LABELS
    )


def _fill_named_envelope_template(document, record, remove_text_boxes=False):
    """Fill envelopes whose labels match the front-end field names."""
    if remove_text_boxes:
        _remove_docx_text_boxes(document)
    department = _recipient_department_value(record)
    contact = record.get("recipient_contact_snapshot") or ""
    replacements = (
        # Replace the template's default value before inserting other values so a
        # newly inserted word such as「掛號」is never replaced a second time.
        ("掛號", record.get("mail_service") or "掛號"),
        ("郵件種類", record.get("mail_service") or "掛號"),
        # Sender postal labels must be replaced before the shorter recipient label
        # 「郵遞區號」, otherwise only the suffix would be replaced.
        ("寄件人郵遞區號", _sender_display_value(record, "sender_postal_code_snapshot", "postal_code")),
        ("寄件人郵寄區號", _sender_display_value(record, "sender_postal_code_snapshot", "postal_code")),
        ("寄件人名稱", _sender_display_value(record, "sender_name_snapshot", "name")),
        ("寄件人地址", _sender_display_value(record, "sender_address_snapshot", "address")),
        ("寄件聯絡人", _sender_display_value(record, "sender_contact_snapshot", "contact")),
        ("寄件人電話", _sender_display_value(record, "sender_phone_snapshot", "phone")),
        ("郵遞區號", resolved_recipient_postal_code(record)),
        ("收件地址", record.get("recipient_address_snapshot") or ""),
        ("寄送內容", record.get("item_description") or ""),
        ("收件人／單位名稱", record.get("recipient_name_snapshot") or ""),
        ("收件部門：收件聯絡人", _recipient_department_contact(record)),
        ("收件部門", department),
        ("收件聯絡人", contact),
    )
    for paragraph in _docx_content_paragraphs(document):
        for old, new in replacements:
            _replace_docx_run_text(paragraph, old, new)


TEXT_BOX_CLEANUP_TEMPLATES = {"信封(大).docx", "信封(小).docx"}
REGISTERED_MARK_TEMPLATE = "信封(小).docx"


def _set_small_envelope_registered_mark(document, mail_service):
    """The small envelope uses a solid block as its registered-mail mark."""
    replacement = "█" if mail_service == "掛號" else ""
    for paragraph in _docx_content_paragraphs(document):
        _replace_docx_run_text(paragraph, "█", replacement)


def _envelope_template_path(template_name):
    safe_name = Path(template_name).name
    if safe_name != template_name or not safe_name.lower().endswith(".docx"):
        raise ValueError("信封範本名稱不正確。")
    template_path = TEMPLATE_DIR / safe_name
    if not template_path.exists():
        raise FileNotFoundError("信封範本尚未放入 doc_templates。")
    return safe_name, template_path


def _render_envelope_document(record, safe_name, template_path):
    mapping = {
        "{{mail_no}}": record.get("mail_no") or "",
        "{{recipient_postal_code}}": resolved_recipient_postal_code(record),
        "{{recipient_name}}": record.get("recipient_name_snapshot") or "",
        "{{recipient_department}}": _recipient_department_value(record),
        "{{recipient_contact}}": record.get("recipient_contact_snapshot") or "",
        "{{recipient_department_contact}}": _recipient_department_contact(record),
        "{{recipient_address}}": record.get("recipient_address_snapshot") or "",
        "{{sender_name}}": _sender_display_value(record, "sender_name_snapshot", "name"),
        "{{sender_postal_code}}": _sender_display_value(record, "sender_postal_code_snapshot", "postal_code"),
        "{{sender_address}}": _sender_display_value(record, "sender_address_snapshot", "address"),
        "{{sender_contact}}": _sender_display_value(record, "sender_contact_snapshot", "contact"),
        "{{sender_phone}}": _sender_display_value(record, "sender_phone_snapshot", "phone"),
        "{{item_description}}": record.get("item_description") or "",
        "{{mail_service}}": record.get("mail_service") or "掛號",
    }
    document = Document(template_path)
    if _uses_named_envelope_labels(document):
        _fill_named_envelope_template(
            document,
            record,
            remove_text_boxes=safe_name.casefold() in TEXT_BOX_CLEANUP_TEMPLATES,
        )
    if safe_name.casefold() == REGISTERED_MARK_TEMPLATE.casefold():
        _set_small_envelope_registered_mark(document, record.get("mail_service"))
    # Placeholder fields remain supported and may be mixed with the Chinese labels.
    for paragraph in document.paragraphs:
        _replace_docx_text(paragraph, mapping)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_docx_text(paragraph, mapping)
    return document


def generate_envelope(record_id, template_name):
    record = get_mail_record(record_id)
    if not record:
        raise ValueError("找不到寄件資料。")
    safe_name, template_path = _envelope_template_path(template_name)
    document = _render_envelope_document(record, safe_name, template_path)
    output = io.BytesIO()
    document.save(output)
    safe_recipient = re.sub(r'[\\/:*?"<>|]+', "_", record["recipient_name_snapshot"])
    return output.getvalue(), f"{record['mail_no']}_{safe_recipient}_信封.docx"


def generate_envelopes(record_ids, template_name):
    """Create one Word file with one envelope per page, preserving input order."""
    ordered_ids = list(dict.fromkeys(record_ids or []))
    if not ordered_ids:
        raise ValueError("請至少選擇一筆寄件。")

    safe_name, template_path = _envelope_template_path(template_name)
    combined = None
    for record_id in ordered_ids:
        record = get_mail_record(record_id)
        if not record:
            raise ValueError(f"找不到寄件資料（ID：{record_id}）。")
        page = _render_envelope_document(record, safe_name, template_path)
        if combined is None:
            combined = page
            continue

        # All pages use the same template, so styles, relationships and page
        # settings already exist in the first document. Keep its final sectPr
        # and append only the subsequent page's body content. Mark the first
        # paragraph with pageBreakBefore instead of inserting a standalone
        # page-break paragraph: full-page envelope templates can push that
        # standalone break onto a blank page and then skip another page.
        page_elements = [
            deepcopy(element)
            for element in page.element.body
            if element.tag != qn("w:sectPr")
        ]
        first_paragraph = next(
            (
                paragraph
                for element in page_elements
                for paragraph in (
                    [element]
                    if element.tag == qn("w:p")
                    else list(element.iter(qn("w:p")))
                )
            ),
            None,
        )
        if first_paragraph is not None:
            paragraph_properties = first_paragraph.find(qn("w:pPr"))
            if paragraph_properties is None:
                paragraph_properties = OxmlElement("w:pPr")
                first_paragraph.insert(0, paragraph_properties)
            if paragraph_properties.find(qn("w:pageBreakBefore")) is None:
                paragraph_properties.insert(0, OxmlElement("w:pageBreakBefore"))
        for element in page_elements:
            combined.element.body.insert(-1, element)

    output = io.BytesIO()
    combined.save(output)
    return output.getvalue(), f"信封套印_{len(ordered_ids)}件.docx"


def save_receipt(uploaded_file, batch_no=None):
    data = uploaded_file.getvalue()
    if not data:
        raise ValueError("收據檔案是空的。")
    if len(data) > MAX_RECEIPT_BYTES:
        raise ValueError("收據檔案不可超過 15 MB。")
    original_name = Path(uploaded_file.name).name
    extension = Path(original_name).suffix.lower()
    if extension not in ALLOWED_RECEIPTS:
        raise ValueError("只接受 PDF、JPG、JPEG 或 PNG。")
    if extension == ".pdf" and not data.startswith(b"%PDF"):
        raise ValueError("檔案內容不是有效的 PDF。")
    RECEIPT_DIR.mkdir(parents=True, exist_ok=True)
    # 以批次號命名，方便直接在資料夾找檔；同名時加流水號，不覆蓋舊檔。
    base_name = re.sub(r'[^0-9A-Za-z_-]+', "_", batch_no or "") or uuid.uuid4().hex
    stored_name = f"{base_name}{extension}"
    counter = 2
    while (RECEIPT_DIR / stored_name).exists():
        stored_name = f"{base_name}-{counter}{extension}"
        counter += 1
    final_path = RECEIPT_DIR / stored_name
    temporary_path = RECEIPT_DIR / f".{stored_name}.tmp"
    temporary_path.write_bytes(data)
    os.replace(temporary_path, final_path)
    return {
        "path": str(Path("mail_receipts") / stored_name),
        "original_name": original_name,
        "content_type": ALLOWED_RECEIPTS[extension],
    }


def delete_saved_receipt(receipt_meta):
    if not receipt_meta:
        return
    try:
        path = _resolve_receipt_path(receipt_meta["path"])
        if path.exists():
            path.unlink()
    except (ValueError, OSError):
        pass


def _resolve_receipt_path(relative_path):
    candidate = (BASE_PATH / str(relative_path)).resolve()
    root = RECEIPT_DIR.resolve()
    if candidate.parent != root:
        raise ValueError("收據路徑不正確。")
    return candidate


def read_receipt(relative_path):
    if not relative_path:
        return None
    path = _resolve_receipt_path(relative_path)
    return path.read_bytes() if path.exists() else None
